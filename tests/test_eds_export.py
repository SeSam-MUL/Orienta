"""What the EDS export actually writes.

Every number in these tests is hand-checkable against the fixture map, which
is the point: an export is evidence, and a test that only asserts "a file
appeared" would pass just as happily on an export whose areas are off by the
step size squared or whose fractions all use the same denominator.

The fixture is a 10x10 map with a matrix, two chemistry regions (one of them
in two pieces), a dead strip with no EDS signal at all, and four hand-painted
pixels — so the three denominators genuinely differ (100 / 90 / 86) instead of
coinciding and hiding a mix-up between them.
"""
from __future__ import annotations

import csv
import math
import json
import sys
from pathlib import Path

import numpy as np
import pytest
from scipy import ndimage

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.api.services import eds_export as X
from backend.api.services.cif_phase_library import CifPhaseEntry
from backend.api.services.phase_map_store import PhaseMapState

N_ROWS = N_COLS = 10
N_PX = N_ROWS * N_COLS
STEP = 0.5                      # um, square pixels

# Hand-counted from the fixture below. If one of these moves, the fixture
# changed and every expectation in the file has to be re-derived — which is
# the intended tripwire.
N_REGION1 = 13                  # 3x3 block plus a 2x2 block, two pieces
N_REGION2 = 4
N_REGION0 = N_PX - N_REGION1 - N_REGION2     # 83
N_DEAD = 10                     # bottom row: no counts at all
# The whole dead strip lies in region 0, so region 0 is the row where the
# row's own pixel count and its valid-EDS pixel count genuinely disagree.
N_REGION0_VALID = N_REGION0 - N_DEAD         # 73
N_HAND = 4                      # hand-marked unclassified
N_VALID = N_PX - N_DEAD         # 90
N_CLASSIFIED = N_PX - N_DEAD - N_HAND        # 86
N_PARTICLES = 4                 # 1 matrix + 2 pieces of region 1 + 1 region 2


def _entry(name, comp, **kw):
    return CifPhaseEntry(
        key=name, cif_filename=name, formula=kw.get("formula", name),
        space_group=kw.get("space_group", "Fm-3m"),
        space_group_number=kw.get("sgn", 225),
        crystal_system=kw.get("system", "cubic"),
        composition=comp, elements=sorted(comp),
    )


ENTRIES = [
    _entry("Al.cif", {"Al": 100.0}),
    _entry("Si.cif", {"Si": 100.0}),
    # Names copper, which this scan never measured. That is what puts a
    # `not measured` column in every table.
    _entry("Cu2Si.cif", {"Cu": 66.67, "Si": 33.33}),
]


@pytest.fixture
def fixture():
    """``(state, at_maps)`` — the map every expectation in this file counts."""
    region = np.zeros((N_ROWS, N_COLS), dtype=np.int32)
    region[1:4, 1:4] = 1                 # 9 px
    region[6:8, 1:3] = 1                 # 4 px, a SECOND piece of region 1
    region[6:8, 6:8] = 2                 # 4 px

    flat_region = region.ravel()
    al = np.full(N_PX, 95.0)
    si = np.full(N_PX, 5.0)
    al[flat_region == 1], si[flat_region == 1] = 40.0, 60.0
    al[flat_region == 2], si[flat_region == 2] = 45.0, 55.0

    # A dead strip: no counts anywhere, so `has_chemistry` reads False and
    # frac_of_valid_eds has a denominator that is not the whole raster.
    dead = np.zeros(N_PX, dtype=bool)
    dead[(N_ROWS - 1) * N_COLS:] = True
    al[dead] = 0.0
    si[dead] = 0.0

    phase = np.where(flat_region > 0, 1, 0).astype(np.int32)
    phase[dead] = -1

    # Four hand-painted "this is not a phase" pixels, well away from anything.
    locked = np.zeros(N_PX, dtype=bool)
    hand = np.array([40, 41, 50, 51])
    phase[hand] = -1
    locked[hand] = True

    state = PhaseMapState(
        phase_grid=phase.reshape(N_ROWS, N_COLS),
        score_grid=np.full((N_ROWS, N_COLS), 0.8, dtype=np.float32),
        phase_entries=list(ENTRIES),
        n_rows=N_ROWS, n_cols=N_COLS, tolerance=15.0, min_score=0.3,
        locked_mask=locked.reshape(N_ROWS, N_COLS),
        region_grid=region,
        region_phase=[0, 1, 1],
    )
    return state, {"Al": al, "Si": si}


def _geom(step=STEP, path="fixture.h5oina"):
    return X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=step,
                            step_y_um=step, source_path=path)


def _build(fixture, **opt_kw):
    state, at_maps = fixture
    opts = X.ExportOptions(hash_source=False, **opt_kw)
    return X.build_tables(state, at_maps, geometry=_geom(), options=opts)


def _row(table, **match):
    for r in table.rows:
        if all(r.get(k) == v for k, v in match.items()):
            return r
    raise AssertionError(f"no row in {table.name} matching {match}: "
                         f"{[{k: r.get(k) for k in match} for r in table.rows]}")


# ---------------------------------------------------------------------------
# particles: the count has to match an independent labelling
# ---------------------------------------------------------------------------

def test_particle_rows_match_an_independent_scipy_labelling(fixture):
    """Acceptance criterion 1 of the spec, checked the way it is written.

    Not "the exporter agrees with itself" — a second, independent run of
    ``ndimage.label`` over the same masks, summed across regions.
    """
    state, _ = fixture
    tables = _build(fixture)

    struct = ndimage.generate_binary_structure(2, 2)          # 8-connectivity
    expected = 0
    for sid in range(len(state.region_phase)):
        _lab, n = ndimage.label(state.region_grid == sid, structure=struct)
        expected += n

    assert expected == N_PARTICLES
    assert len(tables.particles) == expected


def test_a_region_in_two_pieces_gives_two_particle_rows(fixture):
    """The whole reason particles are not regions.

    Region 1 is 13 px in two disconnected blocks. As a region it is one row
    with one mean; as particles it is 9 px and 4 px, which is what a size
    distribution is made of.
    """
    tables = _build(fixture)
    rows = [r for r in tables.particles.rows if r["region_id"] == 1]
    assert sorted(r["n_px"] for r in rows) == [4, 9]
    assert _row(tables.regions, region_id=1)["n_px"] == N_REGION1


def test_connectivity_four_splits_differently_and_is_recorded(fixture):
    """The particle count is not comparable across the two settings.

    Which is exactly why the value used is in the provenance rather than
    being a constant somebody has to go and read in the source.
    """
    tables = _build(fixture, connectivity=4)
    assert tables.provenance["particles"]["connectivity"] == 4
    assert _build(fixture).provenance["particles"]["connectivity"] == 8


# ---------------------------------------------------------------------------
# areas: exact, or absent
# ---------------------------------------------------------------------------

def test_area_is_exactly_n_px_times_the_two_steps(fixture):
    tables = _build(fixture)
    for r in tables.particles.rows:
        assert r["area_um2"] == pytest.approx(r["n_px"] * STEP * STEP)
    for r in tables.regions.rows:
        assert r["area_um2"] == pytest.approx(r["n_px"] * STEP * STEP)
    for r in tables.phases.rows:
        assert r["area_um2"] == pytest.approx(r["n_px"] * STEP * STEP)


def test_anisotropic_pixels_use_both_steps(fixture):
    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=0.5,
                            step_y_um=0.25, source_path="fixture.h5oina")
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=2)
    assert r["area_um2"] == pytest.approx(N_REGION2 * 0.5 * 0.25)


def test_every_um_column_vanishes_when_the_step_is_unknown(fixture):
    """A guessed scale looks like an answer. There must be no column at all."""
    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=None,
                            step_y_um=None, source_path="fixture.h5oina")
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))

    for table in (tables.phases, tables.regions, tables.particles,
                  tables.definitions):
        offenders = [c for c in table.columns if "_um" in c]
        assert offenders == [], f"{table.name} still offers {offenders}"
        for row in table.rows:
            assert not [k for k in row if "_um" in k]

    step = tables.provenance["step"]
    assert step["area_available"] is False
    assert step["x_um"] is None
    assert "no usable step size" in step["note"]
    assert any(w["code"] == "no_step_size" for w in tables.warnings)


def test_a_zero_step_is_treated_as_unknown_not_as_zero(fixture):
    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=0.0,
                            step_y_um=0.5)
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    assert "area_um2" not in tables.regions.columns


# ---------------------------------------------------------------------------
# "not measured" is not zero
# ---------------------------------------------------------------------------

def test_an_unmeasured_element_reads_not_measured_everywhere(fixture):
    """Cu is named by a participating phase and was never mapped.

    "Cu 0.0 at%" would assert it was looked for and found absent, which is a
    different and false statement.
    """
    tables = _build(fixture)
    for table in (tables.phases, tables.regions, tables.particles):
        assert f"mean_at_pct_Cu" in table.columns
        for row in table.rows:
            assert row["mean_at_pct_Cu"] == X.NOT_MEASURED
            assert row["sd_within_at_pct_Cu"] == X.NOT_MEASURED
    assert "Cu" in tables.provenance["elements"]["not_measured"]
    assert "Cu" not in tables.provenance["elements"]["measured"]


def test_a_measured_element_is_never_the_not_measured_string(fixture):
    tables = _build(fixture)
    for row in tables.regions.rows:
        assert isinstance(row["mean_at_pct_Al"], float)
        assert isinstance(row["mean_at_pct_Si"], float)


def test_the_phase_still_states_its_nominal_copper(fixture):
    """What the phase CLAIMS, next to what was measured.

    The nominal column stays a number even where the measured one cannot be:
    the disagreement between "this phase is two thirds copper" and "copper was
    never mapped" is the thing worth seeing.
    """
    state, at_maps = fixture
    # Give Cu2Si some pixels so it earns a row.
    state.phase_grid[0, 0] = 2
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.phases, phase_name="Cu2Si.cif")
    assert r["nominal_at_pct_Cu"] == pytest.approx(66.67)
    assert r["mean_at_pct_Cu"] == X.NOT_MEASURED


# ---------------------------------------------------------------------------
# the three denominators
# ---------------------------------------------------------------------------

def test_the_three_denominators_are_all_present_and_all_different(fixture):
    tables = _build(fixture)
    counts = tables.provenance["counts"]
    assert counts["n_px_total"] == N_PX
    assert counts["n_px_valid_eds"] == N_VALID
    assert counts["n_px_classified"] == N_CLASSIFIED
    # If these ever coincide the fixture has stopped testing anything.
    assert len({N_PX, N_VALID, N_CLASSIFIED}) == 3


def test_the_particle_exclusion_counts_are_numbers_not_only_prose(fixture):
    """"N found, M below the limit, K touching the edge" must be countable.

    Both counts already existed inside an English warning STRING. A caption
    builder that wanted them had to parse that prose, which is fragile and
    untranslatable, so the frontend correctly refused to. They are numbers now.

    Neither flag removes a row — that is the whole point of the flag-never-
    filter rule — so these counts are what lets a reader state the exclusions
    honestly instead of quietly reporting the survivors.
    """
    # 5, not 4: the fixture's smallest pieces are 4 px, so a limit of 4
    # flags nothing and the test would pass while proving nothing.
    tables = _build(fixture, min_particle_px=5)
    counts = tables.provenance["counts"]
    rows = tables.particles.rows

    below = sum(1 for r in rows if r["below_size_limit"])
    edge = sum(1 for r in rows if r["touches_edge"])

    assert counts["n_particles_below_size_limit"] == below
    assert counts["n_particles_touching_edge"] == edge

    # The fixture has to actually exercise both, or this proves nothing.
    assert below > 0, "fixture has no sub-threshold particle to count"
    assert edge > 0, "fixture has no edge-touching particle to count"

    # Flagged, never removed: every particle is still a row.
    assert counts["n_particles"] == len(rows)
    assert below < len(rows), "if everything is flagged the count is not a filter test"


def test_the_provenance_scalar_intersects_the_two_sets_like_the_tables_do():
    """`frac_classified_of_valid_eds` must not mix two different sets either.

    The per-row columns were fixed to intersect; this summary scalar sat on the
    same defect one level up, and would have contradicted the very tables it
    summarises.

    The shared fixture cannot show it: there every hand-painted pixel is left
    unclassified, so "classified" is a subset of "valid EDS" by accident. The
    divergence needs a pixel that is CLASSIFIED but carries no EDS signal —
    which is exactly what happens when a user paints a phase onto a hole, a
    crack or a shadowed area they can identify by eye but the detector cannot.
    """
    n_rows = n_cols = 10
    n_px = n_rows * n_cols
    n_dead = 10                       # bottom row, no counts at all

    al = np.full(n_px, 95.0)
    si = np.full(n_px, 5.0)
    dead = np.zeros(n_px, dtype=bool)
    dead[(n_rows - 1) * n_cols:] = True
    al[dead] = 0.0
    si[dead] = 0.0

    phase = np.zeros(n_px, dtype=np.int32)      # everything is Al...
    phase[dead] = -1                            # ...except the dead strip
    locked = np.zeros(n_px, dtype=bool)

    # The user paints TWO dead pixels as a real phase.
    painted = np.array([(n_rows - 1) * n_cols, (n_rows - 1) * n_cols + 1])
    phase[painted] = 1
    locked[painted] = True

    state = PhaseMapState(
        phase_grid=phase.reshape(n_rows, n_cols),
        score_grid=np.full((n_rows, n_cols), 0.8, dtype=np.float32),
        phase_entries=list(ENTRIES),
        n_rows=n_rows, n_cols=n_cols, tolerance=15.0, min_score=0.3,
        locked_mask=locked.reshape(n_rows, n_cols),
    )
    counts = _build((state, {"Al": al, "Si": si})).provenance["counts"]

    n_valid = n_px - n_dead                     # 90
    n_classified = n_px - n_dead + len(painted)  # 92 — MORE than are valid
    assert counts["n_px_valid_eds"] == n_valid
    assert counts["n_px_classified"] == n_classified

    n_both = counts["n_px_classified_and_valid_eds"]
    assert n_both == n_valid            # the painted pair is excluded
    assert n_both < counts["n_px_classified"]

    frac = counts["frac_classified_of_valid_eds"]
    assert frac == pytest.approx(n_both / n_valid) == pytest.approx(1.0)
    # The pre-fix value was 92/90 — a "fraction" above 1. Pinned so a revert
    # is caught rather than merely looking plausible.
    assert n_classified / n_valid > 1.0
    assert frac <= 1.0


def test_each_fraction_equals_its_count_over_its_named_denominator(fixture):
    tables = _build(fixture)
    seen_a_difference = False
    for table in (tables.regions, tables.particles):
        for r in table.rows:
            n, n_v = r["n_px"], r["n_px_valid_eds"]
            assert r["frac_of_scan"] == pytest.approx(n / N_PX)
            # NOT n / N_VALID: the numerator is drawn from the same set as
            # the denominator, so it counts only this row's valid-EDS pixels.
            assert r["frac_of_valid_eds"] == pytest.approx(n_v / N_VALID)
            assert r["frac_of_classified"] == pytest.approx(n / N_CLASSIFIED)
            seen_a_difference |= n_v != n
    # If the two never disagree the fixture has stopped testing the point.
    assert seen_a_difference


def test_frac_of_valid_eds_counts_only_the_rows_own_valid_pixels(fixture):
    """The numerator is the intersection, hand-counted off the fixture."""
    tables = _build(fixture)
    r0 = _row(tables.regions, region_id=0)
    assert r0["n_px"] == N_REGION0                       # 83
    assert r0["n_px_valid_eds"] == N_REGION0_VALID       # 73, the dead strip
    assert r0["frac_of_valid_eds"] == pytest.approx(N_REGION0_VALID / N_VALID)
    # The old, wrong value. Spelled out so re-aliasing the two is a failure
    # and not a silent revert.
    assert r0["frac_of_valid_eds"] != pytest.approx(N_REGION0 / N_VALID)


def test_frac_of_valid_eds_closes_to_one_across_the_phase_rows(fixture):
    """The whole reason the numerator had to change.

    Every valid-EDS pixel lands in exactly one phase row (``unclassified``
    included), so the column sums to 1 exactly. With the row total as the
    numerator it summed to N_PX / N_VALID = 100/90, which is not a fraction
    of anything.
    """
    tables = _build(fixture)
    rows = tables.phases.rows
    assert sum(r["n_px_valid_eds"] for r in rows) == N_VALID
    assert sum(r["frac_of_valid_eds"] for r in rows) == pytest.approx(1.0)
    # And it is NOT frac_of_scan wearing a different name.
    assert sum(r["frac_of_scan"] for r in rows) == pytest.approx(1.0)
    assert any(r["frac_of_valid_eds"] != pytest.approx(r["frac_of_scan"])
               for r in rows)


def test_the_unclassified_row_is_where_the_two_fractions_part_company(fixture):
    """It owns the dead strip, which carries no EDS measurement at all."""
    r = _row(_build(fixture).phases, phase_index=-1)
    assert r["n_px"] == N_DEAD + N_HAND                  # 14
    assert r["n_px_valid_eds"] == N_HAND                 # 4 — the dead 10 are out
    assert r["frac_of_valid_eds"] == pytest.approx(N_HAND / N_VALID)
    assert r["frac_of_scan"] == pytest.approx((N_DEAD + N_HAND) / N_PX)


def test_no_row_anywhere_claims_more_valid_pixels_than_it_owns(fixture):
    tables = _build(fixture)
    for table in (tables.phases, tables.regions, tables.particles):
        for r in table.rows:
            assert 0 <= r["n_px_valid_eds"] <= r["n_px"]
            assert r["frac_of_valid_eds"] <= 1.0


def test_no_column_is_called_area_pct(fixture):
    """"I will reject any column whose denominator is not in its own header." """
    tables = _build(fixture)
    for table in (tables.phases, tables.regions, tables.particles,
                  tables.definitions):
        for col in table.columns:
            assert col not in ("area_pct", "percentage", "pct")
            if col.startswith("frac"):
                assert col in ("frac_of_scan", "frac_of_valid_eds",
                               "frac_of_classified"), col


def test_every_fraction_has_its_count_beside_it(fixture):
    """So a binomial interval is recoverable without going back to the file."""
    tables = _build(fixture)
    for table in (tables.phases, tables.regions, tables.particles):
        assert "n_px" in table.columns
        # frac_of_valid_eds has a numerator of its own, so it needs a count of
        # its own — immediately before it, where the reader will look.
        assert "n_px_valid_eds" in table.columns
        assert (table.columns.index("n_px_valid_eds") + 1
                == table.columns.index("frac_of_valid_eds"))


def test_the_denominators_are_spelled_out_in_the_provenance(fixture):
    d = _build(fixture).provenance["denominators"]
    assert "usable EDS" in d["frac_of_valid_eds"]
    assert "received a phase" in d["frac_of_classified"]
    # The intersection is stated, not left to be inferred from the name.
    assert "n_px_valid_eds" in d["frac_of_valid_eds"]
    assert "SAME SET" in d["frac_of_valid_eds"]


# ---------------------------------------------------------------------------
# unclassified is a row
# ---------------------------------------------------------------------------

def test_unclassified_is_a_real_row_with_its_own_area_and_composition(fixture):
    tables = _build(fixture)
    r = _row(tables.phases, phase_index=-1)
    assert r["phase_name"] == "unclassified"
    assert r["n_px"] == N_DEAD + N_HAND
    assert r["area_um2"] == pytest.approx((N_DEAD + N_HAND) * STEP * STEP)
    # It has a composition — that composition is the most diagnostic thing in
    # a bad map, and it is why this is a row and not a subtraction from 100 %.
    assert isinstance(r["mean_at_pct_Al"], float)


def test_the_unclassified_row_does_not_divide_itself_by_the_classified_set(fixture):
    """It is by definition not part of that set; the cell is empty, not > 1."""
    tables = _build(fixture)
    r = _row(tables.phases, phase_index=-1)
    assert r["frac_of_classified"] is None
    assert r["frac_of_scan"] == pytest.approx((N_DEAD + N_HAND) / N_PX)


def test_phase_fractions_and_the_unclassified_row_account_for_the_whole_scan(fixture):
    tables = _build(fixture)
    assert sum(r["n_px"] for r in tables.phases.rows) == N_PX


# ---------------------------------------------------------------------------
# provenance / hand edits
# ---------------------------------------------------------------------------

def test_hand_painted_pixels_are_counted_and_graded_not_asserted(fixture):
    """4 painted pixels in an 83-pixel region is ``mixed``, not ``hand``.

    "A boolean that reads as 'a human changed this'" has to be proportionate.
    Region 0 owns 83 pixels and 4 of them were painted; calling the whole
    region ``hand`` claims a human decided all 83.
    """
    tables = _build(fixture)
    assert tables.provenance["hand_edits"]["pixels_painted"] == N_HAND
    r = _row(tables.regions, region_id=0)
    assert r["n_px_hand"] == N_HAND
    assert r["n_px"] == N_REGION0
    assert r["assignment_source"] == "mixed"
    # The regions nobody touched must NOT claim to have been touched.
    assert _row(tables.regions, region_id=1)["assignment_source"] == "auto"


def test_a_fully_painted_region_says_hand(fixture):
    """The other end of the grading: every pixel painted really is ``hand``."""
    state, at_maps = fixture
    locked = np.zeros(N_PX, dtype=bool)
    locked[state.region_grid.ravel() == 2] = True       # all 4 px of region 2
    state.locked_mask = locked.reshape(N_ROWS, N_COLS)
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=2)
    assert r["n_px_hand"] == r["n_px"] == N_REGION2
    assert r["assignment_source"] == "hand"


def test_the_grader_covers_every_case_it_documents():
    """The five values, straight at the function, so each is pinned."""
    g = X._grade_source
    assert g(n_px=10, n_hand=10, agrees=True) == "hand"
    assert g(n_px=10, n_hand=1, agrees=True) == "mixed"
    assert g(n_px=10, n_hand=0, agrees=True, definition=True) == "definition"
    assert g(n_px=10, n_hand=0, agrees=True) == "auto"
    assert g(n_px=10, n_hand=0, agrees=False) == "untracked"
    # No replay, no verdict: it cannot be graded either way, and "auto" would
    # be a claim rather than a measurement.
    assert g(n_px=10, n_hand=0, agrees=None) == "untracked"


def test_a_region_the_map_moved_without_a_trace_says_untracked_not_auto(fixture):
    """Region -> phase naming leaves no trace, so ``auto`` would be a claim.

    ``assign_region_phase`` deliberately does not lock, so a region renamed
    by hand is indistinguishable in the stored map from one the classifier
    named — except that the chemistry no longer agrees with it. That is the
    one signal available, and it must produce ``untracked`` rather than a
    positive assertion of ``auto``.
    """
    state, at_maps = fixture
    state.region_phase[1] = 0                     # user says Al, chemistry says Si
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=1)
    assert r["n_px_hand"] == 0
    assert r["phase_agrees"] is False
    assert r["assignment_source"] == "untracked"


def test_untracked_hand_edits_are_null_and_say_so_rather_than_claiming_zero(fixture):
    """A merge leaves no trace in the stored map. "0" would be a claim."""
    h = _build(fixture).provenance["hand_edits"]
    assert h["regions_named"] is None
    assert h["merges"] is None
    assert "not tracked" in h["note"]


def test_provenance_records_resolved_values_not_only_requested_ones(fixture):
    """Acceptance criterion 8: ``regions: auto -> 3``, ``matrix: auto -> Al``."""
    tables = _build(fixture)
    c = tables.provenance["classification"]
    assert c["n_clusters_requested"] is None          # nobody pinned a count
    assert c["n_clusters_resolved"] == 3              # and this is what it became
    assert c["matrix_element_requested"] is None
    assert c["matrix_element_resolved"] == "Al"       # measured, not assumed
    assert tables.provenance["step"]["x_um"] == STEP


def test_a_requested_value_is_echoed_beside_its_resolved_one(fixture):
    tables = _build(fixture, requested={"n_clusters": 8, "mode": "cluster"})
    c = tables.provenance["classification"]
    assert c["n_clusters_requested"] == 8
    assert c["n_clusters_resolved"] == 3
    assert c["mode"] == "cluster"


def test_the_inert_tolerance_slider_is_recorded_as_inert(fixture):
    """Spec section 9. A user who tunes it and sees nothing must be told why."""
    c = _build(fixture).provenance["classification"]
    assert c["tolerance"] == 15.0
    assert c["tolerance_effective"] is False
    assert "inert" in c["tolerance_note"]


def test_provenance_carries_the_determinism_statement(fixture):
    d = _build(fixture).provenance["determinism"]
    assert d["random_state"] == 0
    assert d["n_init"] == 10
    assert d["deterministic"] is True


def test_provenance_carries_the_caveat_as_literal_text(fixture):
    caveat = _build(fixture).provenance["caveat"]
    assert "SEMI-QUANTITATIVE" in caveat
    assert "no full ZAF correction" in caveat


def test_provenance_names_the_elements_dropped_from_scoring(fixture):
    els = _build(fixture).provenance["elements"]
    assert set(els["excluded_from_scoring"]) == {"C", "O"}


def test_the_excluded_note_no_longer_makes_a_false_claim(fixture):
    """It said C and O were dropped from the composition renormalisation.

    Measured on the real SampleB export, the per-pixel at% sum INCLUDING C
    and O is 100.000000 everywhere and EXCLUDING them runs 91.297860 to
    100.000000. They are in the renormalisation; they are out of the
    SCORING. A reader who believed the old sentence read every at% cell
    against the wrong denominator.
    """
    els = _build(fixture).provenance["elements"]
    note = els["excluded_note"]
    assert "ARE included in the reported at% columns" in note
    assert "SCORING" in note
    # The old key survives, empty, saying why — deleting it would leave a
    # reader who pinned it with no signal at all.
    assert els["excluded_from_renormalisation"] == []
    assert "nothing is excluded" in els["excluded_from_renormalisation_note"]


def test_provenance_carries_the_ambiguity_threshold_and_the_crofton_note(fixture):
    prov = _build(fixture).provenance
    assert prov["classification"]["ambiguity_threshold_at_pct"] == pytest.approx(0.01)
    assert "Crofton" in prov["particles"]["perimeter_note"]


def test_provenance_lists_every_participating_phase_not_only_the_winners(fixture):
    """Cu2Si won no pixels and has no row in phases.csv — it is still a phase
    that was allowed to compete, and "which candidates were on the list" is a
    different question from "which ones won"."""
    tables = _build(fixture)
    names = {p["cif_filename"] for p in tables.provenance["phases_participating"]}
    assert names == {"Al.cif", "Si.cif", "Cu2Si.cif"}
    assert "Cu2Si.cif" not in {r["phase_name"] for r in tables.phases.rows}


def test_source_is_hashed_and_sized(tmp_path, fixture):
    state, at_maps = fixture
    src = tmp_path / "scan.h5oina"
    src.write_bytes(b"not a real scan, but a real sha256")
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=STEP,
                            step_y_um=STEP, source_path=str(src))
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=True))
    import hashlib
    assert tables.provenance["source"]["sha256"] == \
        hashlib.sha256(src.read_bytes()).hexdigest()
    assert tables.provenance["source"]["bytes"] == src.stat().st_size


def test_skipping_the_hash_says_so_rather_than_omitting_the_field(tmp_path, fixture):
    state, at_maps = fixture
    src = tmp_path / "scan.h5oina"
    src.write_bytes(b"x")
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=STEP,
                            step_y_um=STEP, source_path=str(src))
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    s = tables.provenance["source"]
    assert s["sha256"] is None
    assert "disabled" in s["sha256_note"]


# ---------------------------------------------------------------------------
# margins and the auto-vs-final disagreement
# ---------------------------------------------------------------------------

def test_margin_follows_the_rankings_own_sign_convention(fixture):
    """Distance ranking counts up from the winner; score ranking counts down."""
    ranked = [(0, 1.5), (1, 4.0), (2, 9.0)]
    assert X._margin(ranked) == pytest.approx(2.5)
    scored = [(0, 0.9), (1, 0.4), (2, 0.1)]
    assert X._margin(scored, higher_is_better=True) == pytest.approx(0.5)
    # Both directions agree that an unopposed winner has no margin.
    assert X._margin([(0, 1.5)]) is None
    assert X._margin([(0, 1.5)], higher_is_better=True) is None
    assert X._margin([]) is None


def test_every_region_with_a_winner_carries_the_deciding_margin(fixture):
    """2.7 against 9.1 is a decision; 2.7 against 2.8 is a coin toss."""
    tables = _build(fixture)
    for r in tables.regions.rows:
        if r["phase_auto_index"] >= 0:
            assert r["margin_score"] is not None
            assert r["score_auto"] >= r["runner_up_score"]
            assert r["margin_score"] == pytest.approx(
                r["score_auto"] - r["runner_up_score"])


def test_the_mad_columns_are_internally_consistent_too(fixture):
    tables = _build(fixture)
    for r in tables.regions.rows:
        if r["gap_mad_at_pct"] is None:
            continue
        assert r["margin_mad_at_pct"] is not None
        assert r["margin_mad_at_pct"] >= 0.0


def test_a_single_candidate_library_has_no_runner_up_and_says_so(fixture):
    """The one path that legitimately exports a winner without a margin.

    The old docstring claimed the invariant "never export a winner without a
    margin" and four paths broke it. This is the one that is not a bug: an
    unopposed winner is a fact worth seeing, not a confident one.
    """
    state, at_maps = fixture
    state.phase_entries = [ENTRIES[0]]
    state.region_phase = [0, 0, 0]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    # Region 0 is the 95 at% Al matrix, so the lone candidate genuinely wins
    # rather than merely being the only name on the list.
    r = _row(tables.regions, region_id=0)
    assert r["phase_auto"] == "Al.cif"
    assert r["score_auto"] > 0.3
    assert r["runner_up_phase"] == ""
    assert r["runner_up_score"] is None
    assert r["margin_score"] is None


def test_a_region_no_candidate_clears_the_floor_for_names_none_of_them(fixture):
    """``min_score`` is the classifier's own floor and the replay honours it.

    Region 1 is 60 at% silicon; hand it a library containing only pure
    aluminium and the honest answer is "no phase", which is exactly what the
    classifier would do. ``score_auto`` is still reported — "the best
    chemistry could manage was below the floor" is the fact a reader needs.
    """
    state, at_maps = fixture
    state.phase_entries = [ENTRIES[0]]
    state.region_phase = [0, -1, -1]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=1)
    assert r["phase_auto"] == ""
    assert r["phase_auto_index"] == -1
    assert r["score_auto"] is not None and r["score_auto"] < 0.3
    # -1 == -1 is an AGREEMENT. The old code called this False.
    assert r["phase_agrees"] is True


def test_a_zero_pixel_region_emits_no_winner_at_all(fixture):
    """The second path that used to write phase_final with a null margin.

    It no longer writes an auto winner either, so there is no winner left
    unaccompanied — the row is empty on both sides of the comparison and
    ``phase_agrees`` is null rather than a fabricated verdict.
    """
    state, at_maps = fixture
    state.region_phase = list(state.region_phase) + [1]   # region 3 owns nothing
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=3)
    assert r["n_px"] == 0
    assert r["phase_auto"] == ""
    assert r["phase_auto_index"] == -1
    assert r["score_auto"] is None
    assert r["margin_score"] is None
    assert r["phase_agrees"] is None


def test_phase_auto_comes_from_the_scorer_that_made_the_map(fixture):
    """THE anti-drift test for defect 1, and the one to break first.

    ``phase_auto`` must be the answer ``eds_clustering._match_labels`` would
    give — ``chemistry_score.score_phase_ratio`` on the region's eroded
    interior — and not a second metric wearing its name. This runs the real
    scorer against the real ``_interior_mask`` and demands the exported cell
    agree bit for bit, so a future edit that quietly swaps the metric back
    turns this red instead of turning 4 of 7 regions of a real scan into a
    fabricated manual-override claim.
    """
    from backend.api.services.chemistry_score import (
        background_levels, infer_matrix_element, score_phase_ratio)
    from backend.api.services.eds_clustering import _interior_mask

    state, at_maps = fixture
    tables = _build(fixture)
    flat_region = state.region_grid.ravel()
    bg = background_levels(at_maps)
    mx = infer_matrix_element(at_maps)

    for sid in range(len(state.region_phase)):
        # The reference: the classification's own erosion, its own scorer,
        # its own no_data_score, its own map-level background.
        interior = _interior_mask(flat_region, sid, N_ROWS, N_COLS)
        one = {el: np.array([float(np.asarray(v, float).ravel()[interior].mean())])
               for el, v in at_maps.items()}
        scored = sorted(
            ((i, float(score_phase_ratio(one, e.composition, matrix_element=mx,
                                         no_data_score=0.0, background=bg)[0]))
             for i, e in enumerate(state.phase_entries)),
            key=lambda t: (-t[1], t[0]))

        r = _row(tables.regions, region_id=sid)
        best_i, best_s = scored[0]
        assert r["score_auto"] == pytest.approx(best_s, abs=1e-9), f"region {sid}"
        assert r["runner_up_score"] == pytest.approx(scored[1][1], abs=1e-9)
        want = X._phase_name(state.phase_entries[best_i]) if best_s >= state.min_score else ""
        assert r["phase_auto"] == want, f"region {sid}"


def test_the_cropped_erosion_is_the_classifications_erosion():
    """``_interior_flat`` is a faster route to ``_interior_mask``, not another one.

    It crops to the bounding box and pads with one ring of False, which is
    ``border_value=0``. Checked against the classification's own function on
    the cases where a cropping bug would hide: objects touching each raster
    edge, a corner, a single pixel, a one-pixel-wide line, and a solid block
    that does have an interior. "A shade different from the classification"
    is the defect class this whole file is being repaired for, so this is
    exact equality and not an approximation.
    """
    from backend.api.services.eds_clustering import _interior_mask

    shapes = {
        "top edge": (np.s_[0:3, 4:8],),
        "bottom edge": (np.s_[N_ROWS - 3:N_ROWS, 4:8],),
        "left edge": (np.s_[3:7, 0:3],),
        "right edge": (np.s_[3:7, N_COLS - 3:N_COLS],),
        "corner": (np.s_[0:2, 0:2],),
        "single pixel": (np.s_[5:6, 5:6],),
        "thin line": (np.s_[2:3, 1:9],),
        "solid block": (np.s_[3:8, 3:8],),
        "two pieces": (np.s_[1:4, 1:4], np.s_[6:9, 6:9]),
        "whole raster": (np.s_[0:N_ROWS, 0:N_COLS],),
    }
    for name, slices in shapes.items():
        grid = np.zeros((N_ROWS, N_COLS), dtype=np.int32)
        for sl in slices:
            grid[sl] = 1
        mine = X._interior_flat(grid == 1)
        theirs = _interior_mask(grid.ravel(), 1, N_ROWS, N_COLS)
        assert np.array_equal(mine, theirs), name
    # An empty mask has no interior and no fallback to fall back to.
    empty = np.zeros((N_ROWS, N_COLS), dtype=bool)
    assert not X._interior_flat(empty).any()


def test_phase_auto_is_not_the_mad_ranking_wearing_its_name(fixture):
    """The two metrics genuinely differ, and the columns keep them apart.

    A test that only checked "phase_auto is filled in" would have passed
    happily on the broken export. This one demands the two rankings be
    reported in their own units under their own names: a score is 0..1 and
    higher is better, a MAD is an at% distance and lower is better, so the
    same cell cannot honestly hold both.
    """
    tables = _build(fixture)
    r = _row(tables.regions, region_id=1)
    assert 0.0 <= r["score_auto"] <= 1.0
    assert r["gap_mad_at_pct"] > 1.0            # an at% distance, not a score
    assert r["phase_second_opinion_mad"]        # named separately, always present
    # And the old names, which held MAD numbers while claiming to be the
    # software's answer, are gone from every table.
    for table in (tables.regions, tables.particles):
        assert "gap_at_pct" not in table.columns
        assert "runner_up_gap_at_pct" not in table.columns
        assert "margin_at_pct" not in table.columns


def test_a_map_nobody_touched_reports_no_disagreement_anywhere(fixture):
    """The user-visible symptom, stated as a test.

    On the real 90x120 SampleB scan the broken export marked 4 of 7 regions
    ``phase_agrees = False`` with not one hand-painted pixel on the map, and
    a tester wrote a Methods paragraph claiming a manual override that never
    happened. The fixture is the same shape of thing: strip the hand edits
    and leave the classifier's own answers in place, and every region must
    agree with itself.
    """
    state, at_maps = fixture
    state.locked_mask = None
    # What the classifier itself would have written into region_phase.
    tables = _build((state, at_maps))
    for r in tables.regions.rows:
        state.region_phase[r["region_id"]] = r["phase_auto_index"]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    disagreed = [r["region_id"] for r in tables.regions.rows
                 if r["phase_agrees"] is False]
    assert disagreed == []
    assert {r["assignment_source"] for r in tables.regions.rows} == {"auto"}


def test_a_phase_rule_gates_the_replay_exactly_as_it_gated_the_map(fixture):
    """Rules decide eligibility, and the replay has to know that.

    The rules travel with the map in ``settings["rules"]``. A blocked phase
    scores 0 and so loses even to a badly-scoring phase that is allowed —
    ``phase_auto`` has to reflect that, or the export reports a winner the
    classifier was forbidden to pick.
    """
    state, at_maps = fixture
    ungated = _row(_build(fixture).regions, region_id=1)
    assert ungated["phase_auto"] == "Si.cif"

    state.settings = {
        "rules": {
            "rules": [{
                "phase_key": "Si.cif",
                # Region 1 measures 60 at% Si. This can never be met.
                "elements": [{"element": "Si", "min_at_pct": 99.0}],
            }],
        },
    }
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=1)
    assert r["phase_auto"] != "Si.cif"
    assert tables.provenance["ranking"]["deciding"]["rules_applied"] is True


def test_rank_candidates_reproduces_the_inspectors_own_ranking(monkeypatch, fixture):
    """The anti-drift test.

    The inspector tells a user "AlFeMnSi, 2.7 at% off" and the export has to
    agree to the digit, or the page and the file are using two different
    notions of "close". This runs ``routes/eds.py::_region_detail`` — the
    original — and checks the exporter against ITS numbers rather than against
    a copy of its formula.
    """
    from backend.api.routes import eds as eds_routes

    state, at_maps = fixture
    monkeypatch.setattr(
        eds_routes, "_build_at_pct_maps_for_loaded_file",
        lambda: (at_maps, N_ROWS, N_COLS, "fixture.h5oina"))

    for sid in (0, 1, 2):
        detail = eds_routes._region_detail(state, sid)
        mask = state.region_grid.ravel() == sid
        means = X._ranking_means(at_maps, mask, X._background_levels(at_maps))
        mine = X.rank_candidates(means, state.phase_entries)

        theirs = [(c["phase_index"], c["gap_at_pct"]) for c in detail["candidates"]]
        assert [i for i, _ in mine] == [i for i, _ in theirs], f"region {sid}"
        for (_i, g_mine), (_j, g_theirs) in zip(mine, theirs):
            assert round(g_mine, 2) == pytest.approx(g_theirs)


def test_a_particles_margin_is_its_own_and_not_its_regions(fixture):
    """Defect 2: two particles of different sizes printed one number.

    Region 1 is in two pieces — a 9-pixel block and a 4-pixel block. They
    used to carry the region's 13-pixel margin, identically, and "margin
    0.66 at%" set beside a 4-pixel object is read as describing that object.
    Each piece is now ranked on its own composition; the region's number is
    still there, under a name that says whose it is.
    """
    state, at_maps = fixture
    # Make the two pieces chemically DIFFERENT, so a shared margin is
    # provably wrong rather than merely unjustified.
    flat_region = state.region_grid.ravel()
    piece = np.zeros(N_PX, dtype=bool)
    piece[(np.arange(N_PX) // N_COLS >= 6) & (flat_region == 1)] = True
    at_maps = {"Al": at_maps["Al"].copy(), "Si": at_maps["Si"].copy()}
    at_maps["Al"][piece], at_maps["Si"][piece] = 10.0, 90.0

    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    pieces = [r for r in tables.particles.rows if r["region_id"] == 1]
    assert len(pieces) == 2
    a, b = pieces
    assert a["n_px"] != b["n_px"]
    assert a["margin_score"] != b["margin_score"]
    assert a["score_auto"] != b["score_auto"]
    # The region's own number is present and explicitly labelled, and it is
    # the SAME for both — which is exactly why it needed its own name.
    assert a["region_margin_score"] == b["region_margin_score"]
    assert "region_margin_score" in tables.particles.columns


def test_a_particle_is_ranked_on_the_same_population_a_region_is(fixture):
    """Eroded interior, falling back to the whole object.

    One rule, imported from the classification, so a particle and a region
    of identical pixels cannot be scored on different populations. Region 2
    is a single 2x2 block and is therefore also exactly one particle.
    """
    tables = _build(fixture)
    region = _row(tables.regions, region_id=2)
    particle = _row(tables.particles, region_id=2)
    assert particle["n_px"] == region["n_px"] == N_REGION2
    assert particle["score_auto"] == pytest.approx(region["score_auto"])
    assert particle["margin_score"] == pytest.approx(region["margin_score"])


def test_a_phase_basis_particle_is_ranked_too_rather_than_left_empty(fixture):
    """The fourth path that used to leave a winner marginless.

    On a per-pixel map every particle got ``margin_at_pct = None`` because
    there was no region to inherit from. There is no need to inherit: the
    particle's own composition is measured either way.
    """
    state, at_maps = fixture
    state.region_grid = None
    state.region_phase = []
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert tables.provenance["particles"]["basis"] == "phase"
    scored = [r for r in tables.particles.rows if r["score_auto"] is not None]
    assert scored, "every particle should carry its own ranking"
    for r in scored:
        assert r["margin_score"] is not None
        assert r["phase_agrees"] in (True, False)
        assert r["region_margin_score"] is None


def test_both_sides_of_the_mad_subtraction_are_on_one_basis(fixture):
    """Defect 3: different denominators either side of a minus sign.

    A CIF's nominal composition is C/O-free and sums to 100 by construction.
    A measured at% includes C and O and, with them dropped, sums to less —
    91.30 to 100.00 per pixel on the real SampleB scan. Subtracting one from
    the other charged every phase for carbon it could not contain.
    """
    means = {"Al": 45.0, "Si": 45.0}          # sums to 90, as a C/O-bearing map does
    same = X.to_scored_basis(means)
    assert sum(same.values()) == pytest.approx(100.0)
    assert same["Al"] == pytest.approx(50.0)
    # Nothing to renormalise against is left alone rather than divided by zero.
    assert X.to_scored_basis({"Al": 0.0}) == {"Al": 0.0}
    assert X.to_scored_basis({}) == {}


def test_the_exported_mad_uses_the_renormalised_basis(fixture):
    """And the file says so, and says where it differs from the screen."""
    state, at_maps = fixture
    tables = _build(fixture)
    mask = state.region_grid.ravel() == 1
    raw_means = X._ranking_means(at_maps, mask, X._background_levels(at_maps))
    same = X.rank_candidates(X.to_scored_basis(raw_means), state.phase_entries)

    r = _row(tables.regions, region_id=1)
    assert r["gap_mad_at_pct"] == pytest.approx(same[0][1])
    assert r["margin_mad_at_pct"] == pytest.approx(X._margin(same))

    prov = tables.provenance["ranking"]["second_opinion"]
    assert "RENORMALISED" in prov["composition_basis"]
    assert prov["differs_from_screen"] is True
    assert "rank_candidates" in prov["differs_from_screen_note"]


def test_a_hand_moved_region_shows_both_answers_and_flags_the_disagreement(fixture):
    """"Never overwrite the software's answer."

    Region 1 is 60 at% silicon; the chemistry's own best candidate is Si.cif.
    Point the map at Al.cif instead, as a hand assignment would, and both
    answers have to survive into the file.
    """
    state, at_maps = fixture
    state.region_phase[1] = 0                     # user says Al
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=1)
    assert r["phase_final"] == "Al.cif"
    assert r["phase_auto"] == "Si.cif"
    assert r["phase_agrees"] is False


def test_where_they_agree_the_flag_says_so(fixture):
    tables = _build(fixture)
    r = _row(tables.regions, region_id=1)
    assert r["phase_final"] == "Si.cif"
    assert r["phase_auto"] == "Si.cif"
    assert r["phase_agrees"] is True


# ---------------------------------------------------------------------------
# composition: basis, spread, core
# ---------------------------------------------------------------------------

def test_composition_carries_a_basis_flag(fixture):
    tables = _build(fixture)
    for table in (tables.phases, tables.regions, tables.particles):
        assert "basis" in table.columns
        assert {r["basis"] for r in table.rows} <= {"raw"}
    assert _build(fixture).provenance["composition"]["basis"] == "raw"


def test_the_smoothed_columns_are_omitted_when_the_width_is_unknown(fixture):
    """The store does not persist the smoothing width. Recomputing at a
    default would label a column "smoothed" that was smoothed differently
    from the thing it claims to explain."""
    tables = _build(fixture)
    assert not [c for c in tables.regions.columns if c.startswith("smoothed_")]
    note = tables.provenance["classification"]["smoothing"]["note"]
    assert "not persisted" in note


def test_the_smoothed_columns_appear_when_the_caller_states_the_width(fixture):
    tables = _build(fixture, smoothing_scale_px=3)
    assert "smoothed_mean_at_pct_Si" in tables.regions.columns
    assert tables.provenance["classification"]["smoothing"]["scale_px"] == 3
    r = _row(tables.regions, region_id=1)
    # Smoothing pulls a small particle toward the matrix — the measured
    # difference that makes the basis flag necessary in the first place.
    assert r["smoothed_mean_at_pct_Si"] < r["mean_at_pct_Si"]


def test_the_spread_is_a_population_sd_and_is_named_sd_within(fixture):
    """Not "±", not "uncertainty", and emphatically not a standard error."""
    tables = _build(fixture)
    assert "sd_within_at_pct_Si" in tables.regions.columns
    assert not [c for c in tables.regions.columns
                if "stderr" in c or "sem" in c or "uncertainty" in c]

    # Region 1 is homogeneous; region 0 is matrix plus a dead strip, so its
    # spread is a real number and computable by hand.
    assert _row(tables.regions, region_id=1)["sd_within_at_pct_Si"] == \
        pytest.approx(0.0)
    r0 = _row(tables.regions, region_id=0)
    si0 = np.concatenate([np.full(N_REGION0 - N_DEAD, 5.0),
                          np.zeros(N_DEAD)])
    assert r0["sd_within_at_pct_Si"] == pytest.approx(si0.std())   # ddof=0

    note = tables.provenance["composition"]["spread_note"]
    assert "spatially correlated" in note
    assert "No standard error is exported" in note


def test_particle_composition_is_exported_both_all_pixel_and_core_only(fixture):
    """The rim of a small particle is a particle-plus-matrix mixture.

    On the 3x3 block the core is the single centre pixel, so the two
    compositions are computable by hand and — on this synthetic map, where
    there is no rim mixing — identical. The columns exist so that on real data
    the difference is visible instead of being averaged away.
    """
    tables = _build(fixture)
    big = _row(tables.particles, region_id=1, n_px=9)
    assert big["n_px_core"] == 1
    assert big["mean_at_pct_Si"] == pytest.approx(60.0)
    assert big["core_mean_at_pct_Si"] == pytest.approx(60.0)


def test_a_particle_that_erodes_to_nothing_has_an_empty_core_composition(fixture):
    """``n_px_core == 0`` is information, not an error, and the cell next to
    it must not invent a number."""
    tables = _build(fixture)
    small = _row(tables.particles, region_id=2)
    assert small["n_px"] == 4
    assert small["n_px_core"] == 0
    assert small["core_mean_at_pct_Si"] is None
    assert small["mean_at_pct_Si"] == pytest.approx(55.0)


# ---------------------------------------------------------------------------
# the columns the testers asked for, from data already in hand
# ---------------------------------------------------------------------------

def test_sqrt_area_is_murakamis_parameter_and_is_not_the_ecd(fixture):
    """sqrt(area), the length a failure analyst rates an inclusion by.

    Pinned against ECD as well as against sqrt(area), because the two are
    within 12 % of each other for a compact object and a column that quietly
    duplicated ECD would look right in a spot check.
    """
    tables = _build(fixture)
    for table in (tables.regions, tables.particles, tables.phases):
        for r in table.rows:
            if not r.get("area_um2"):
                continue
            assert r["sqrt_area_um"] == pytest.approx(math.sqrt(r["area_um2"]))
            if "ecd_um" in r and r["ecd_um"]:
                # sqrt(pi)/2 = 0.8862 for any shape, by construction.
                assert r["sqrt_area_um"] / r["ecd_um"] == pytest.approx(
                    math.sqrt(math.pi) / 2.0)


def test_sqrt_area_vanishes_with_every_other_micrometre_column(fixture):
    tables = _build(fixture, )
    no_scale = X.build_tables(
        fixture[0], fixture[1],
        geometry=X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS),
        options=X.ExportOptions(hash_source=False))
    assert "sqrt_area_um" in tables.regions.columns
    for table in (no_scale.regions, no_scale.particles, no_scale.phases):
        assert not [c for c in table.columns if c.endswith("_um")
                    or c.endswith("_um2") or c.endswith("_deg")]


def test_feret_angle_states_its_convention_and_measures_it(fixture):
    """From +x, counter-clockwise, [0, 180). Measured on known shapes.

    A horizontal 1x7 bar and a vertical 7x1 one. Note that ``feret_max`` on a
    rectangle is its DIAGONAL — ``eds_particles.scale_to_um`` says so — so
    the answers are 171.87 and 98.13, not 0 and 90, and this test asserts the
    analytic diagonal rather than the convenient round number.

    Both land ABOVE 90: the hull runs over pixel corners with the row index
    increasing downward, so both diagonals point up-and-left. A convention
    that forgot to negate the row component would answer 8.13 and 81.87 —
    both below 90 — which is the discriminator this pair exists for.
    """
    region = np.zeros((N_ROWS, N_COLS), dtype=np.int32)
    region[2, 1:8] = 1                 # 1 px tall, 7 px wide
    region[1:8, 9] = 2                 # 7 px tall, 1 px wide
    flat = region.ravel()
    al, si = np.full(N_PX, 95.0), np.full(N_PX, 5.0)
    al[flat == 1], si[flat == 1] = 40.0, 60.0
    al[flat == 2], si[flat == 2] = 45.0, 55.0
    state = PhaseMapState(
        phase_grid=np.where(flat > 0, 1, 0).astype(np.int32).reshape(N_ROWS, N_COLS),
        score_grid=np.full((N_ROWS, N_COLS), 0.8, dtype=np.float32),
        phase_entries=list(ENTRIES), n_rows=N_ROWS, n_cols=N_COLS,
        tolerance=15.0, min_score=0.3, region_grid=region,
        region_phase=[0, 1, 1])
    tables = X.build_tables(state, {"Al": al, "Si": si}, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))

    flat_bar = _row(tables.particles, region_id=1)["feret_angle_deg"]
    tall_bar = _row(tables.particles, region_id=2)["feret_angle_deg"]
    assert flat_bar == pytest.approx(math.degrees(math.atan2(-1.0, 7.0)) % 180.0)
    assert tall_bar == pytest.approx(math.degrees(math.atan2(-7.0, 1.0)) % 180.0)
    assert 90.0 < flat_bar < 180.0 and 90.0 < tall_bar < 180.0
    # The wide bar is nearly along +x, the tall one nearly along +y.
    assert flat_bar > 170.0 and tall_bar < 100.0

    note = tables.provenance["particles"]["feret_angle_note"]
    assert "COUNTER-CLOCKWISE" in note and "+X AXIS" in note and "[0, 180)" in note


def test_a_diagonal_feret_angle_is_not_mirrored(fixture):
    """The sign test that a horizontal or a vertical bar cannot make.

    A bar running down-and-right on screen goes DOWN in row index, so
    counter-clockwise-from-+x puts it in the second quadrant, folded to
    135 deg. A convention that forgot the raster's row index increases
    downward would answer 45.
    """
    region = np.zeros((N_ROWS, N_COLS), dtype=np.int32)
    for i in range(7):
        region[1 + i, 1 + i] = 1       # top-left to bottom-right
    flat = region.ravel()
    al, si = np.full(N_PX, 95.0), np.full(N_PX, 5.0)
    al[flat == 1], si[flat == 1] = 40.0, 60.0
    state = PhaseMapState(
        phase_grid=np.where(flat > 0, 1, 0).astype(np.int32).reshape(N_ROWS, N_COLS),
        score_grid=np.full((N_ROWS, N_COLS), 0.8, dtype=np.float32),
        phase_entries=list(ENTRIES), n_rows=N_ROWS, n_cols=N_COLS,
        tolerance=15.0, min_score=0.3, region_grid=region,
        region_phase=[0, 1])
    tables = X.build_tables(state, {"Al": al, "Si": si}, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert _row(tables.particles, region_id=1)["feret_angle_deg"] == pytest.approx(135.0)


def test_enrichment_sits_beside_the_background_it_was_divided_by(fixture):
    """"I asked for enrichment WITH its background; I got neither."

    Pinned against ``eds_wand.selection_stats`` — the same definition the
    classifier's own enrichment gate uses — because this project has already
    shipped an "Al 49.68x enriched" readout from a hand-rolled second
    definition on a map whose background IS aluminium.
    """
    from backend.api.services.eds_wand import selection_stats

    state, at_maps = fixture
    tables = _build(fixture)
    bg = X._background_levels(at_maps)
    mask = state.region_grid.ravel() == 1
    theirs = selection_stats(at_maps, mask, background=bg)["enrichment"]

    r = _row(tables.regions, region_id=1)
    for el in ("Al", "Si"):
        assert r[f"enrichment_{el}"] == pytest.approx(theirs[el], abs=0.005)
        assert r[f"background_at_pct_{el}"] == pytest.approx(bg[el] * 100.0)
    # The whole point: the divisor is in the file, not only in the JSON.
    assert "background_at_pct_Si" in tables.regions.columns
    assert "background_at_pct_Si" in tables.particles.columns
    assert (tables.provenance["elements"]["background_at_pct"]["Si"]
            == pytest.approx(bg["Si"] * 100.0))


def test_an_unmeasured_element_has_no_enrichment_and_no_background(fixture):
    """``not measured``, for the same reason the composition cell says it."""
    r = _row(_build(fixture).regions, region_id=1)
    assert r["enrichment_Cu"] == X.NOT_MEASURED
    assert r["background_at_pct_Cu"] == X.NOT_MEASURED


def test_scored_at_pct_sum_is_what_survives_the_c_and_o_drop(fixture):
    """The recoverable half of the raw-sum request.

    The fixture measures no carbon or oxygen, so the sum is the full 100 —
    which is the right answer and pins the arithmetic. The interesting case
    is the one measured on the real scan (91.30 to 100.00), and the note in
    the provenance is what tells a reader which they are looking at.
    """
    tables = _build(fixture)
    r = _row(tables.regions, region_id=1)
    assert r["scored_at_pct_sum"] == pytest.approx(100.0)
    for table in (tables.regions, tables.particles, tables.phases):
        assert "scored_at_pct_sum" in table.columns
    assert "scored_at_pct_sum" in tables.pixel_columns


def test_carbon_is_in_the_at_pct_and_out_of_the_scored_sum(fixture):
    """The measurement behind the corrected note, as a test.

    Give the map 10 at% carbon everywhere and the at% columns still sum to
    100 across every measured element — carbon included — while
    scored_at_pct_sum drops to 90. That is precisely the distinction the old
    provenance note got backwards.
    """
    state, at_maps = fixture
    al = at_maps["Al"] * 0.9
    si = at_maps["Si"] * 0.9
    c = np.where(at_maps["Al"] + at_maps["Si"] > 0, 10.0, 0.0)
    tables = X.build_tables(state, {"Al": al, "Si": si, "C": c},
                            geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = _row(tables.regions, region_id=1)
    total = sum(r[f"mean_at_pct_{el}"] for el in ("Al", "Si", "C"))
    assert total == pytest.approx(100.0)
    assert r["scored_at_pct_sum"] == pytest.approx(90.0)


def test_the_raw_sum_is_declared_unrecoverable_rather_than_omitted(fixture):
    """"Say so explicitly in provenance rather than omitting it silently."""
    comp = _build(fixture).provenance["composition"]
    assert "NOT RECOVERABLE" in comp["raw_sum_note"]
    assert "at% maps only" in comp["raw_sum_note"]
    assert "scored_at_pct_sum is the recoverable neighbour" in comp["raw_sum_note"]
    assert "91.30 to" in comp["scored_at_pct_sum_note"]


def test_the_crofton_correction_survives_someone_opening_only_the_csv(fixture):
    """A column called ``perimeter_um`` does not say it is an estimate."""
    tables = _build(fixture)
    assert "perimeter_crofton_um" in tables.particles.columns
    assert "perimeter_um" not in tables.particles.columns
    r = tables.particles.rows[0]
    assert r["perimeter_crofton_um"] > 0
    assert "perimeter_crofton_um (renamed" in tables.provenance["particles"]["perimeter_note"]


def test_edge_particles_are_flagged_because_their_size_is_a_lower_bound(fixture):
    tables = _build(fixture)
    matrix = _row(tables.particles, region_id=0)
    assert matrix["touches_edge"] is True
    assert _row(tables.particles, region_id=2)["touches_edge"] is False


# ---------------------------------------------------------------------------
# min_particle_px flags, never deletes
# ---------------------------------------------------------------------------

def test_min_particle_px_flags_without_deleting_a_single_row(fixture):
    """Silent exclusion was named a disqualifier."""
    unfiltered = _build(fixture)
    filtered = _build(fixture, min_particle_px=5)

    assert len(filtered.particles) == len(unfiltered.particles) == N_PARTICLES
    flagged = [r for r in filtered.particles.rows if r["below_size_limit"]]
    assert sorted(r["n_px"] for r in flagged) == [4, 4]
    assert all(r["n_px"] < 5 for r in flagged)
    assert filtered.provenance["particles"]["min_particle_px"] == 5
    assert "flag, never a filter" in \
        filtered.provenance["particles"]["min_particle_px_note"]


def test_the_flagged_particles_are_counted_in_a_warning(fixture):
    tables = _build(fixture, min_particle_px=5)
    w = [x for x in tables.warnings if x["code"] == "below_size_limit"]
    assert len(w) == 1
    assert "2 of 4 particles" in w[0]["message"]


def test_min_particle_px_zero_flags_nothing(fixture):
    tables = _build(fixture)
    assert not any(r["below_size_limit"] for r in tables.particles.rows)


# ---------------------------------------------------------------------------
# locale
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("decimal,delimiter", [
    (".", ","), (".", ";"), (".", "\t"), (",", ";"), (",", "\t"),
])
def test_accepted_locale_combinations_round_trip(tmp_path, fixture,
                                                 decimal, delimiter):
    tables = _build(fixture)
    opts = X.ExportOptions(decimal=decimal, delimiter=delimiter,
                           hash_source=False,
                           artefacts={"phases": True, "regions": False,
                                      "particles": False, "definitions": False,
                                      "xlsx": False})
    res = X.write_export(tables, tmp_path, options=opts)
    path = Path(res.folder) / "phases.csv"
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh, delimiter=delimiter))
    header, first = rows[0], rows[1]
    assert header[:2] == ["phase_index", "phase_name"]
    cell = first[header.index("frac_of_scan")]
    # The value is a real fraction written in the caller's convention, and it
    # has not been split across two fields.
    assert float(cell.replace(",", ".")) == pytest.approx(
        tables.phases.rows[0]["frac_of_scan"])
    assert len(first) == len(header)
    if decimal == ",":
        assert "." not in cell


def test_decimal_comma_with_a_comma_delimiter_is_refused_loudly(fixture):
    """"3,14" in a comma-separated file is two fields, and every column after
    it shifts by one. A German Excel then reads a different number in every
    column, and nothing looks wrong."""
    with pytest.raises(ValueError, match="decimal comma"):
        X.ExportOptions(decimal=",", delimiter=",").validate()


@pytest.mark.parametrize("kw", [
    {"decimal": "'"}, {"delimiter": "|"}, {"connectivity": 6},
    {"min_particle_px": -1}, {"on_existing": "overwrite"},
])
def test_unusable_options_are_refused(kw):
    with pytest.raises(ValueError):
        X.ExportOptions(**kw).validate()


def test_numbers_are_written_at_full_precision(tmp_path, fixture):
    """Rounding is presentation. ``1/3`` has to come back as ``1/3``."""
    tables = _build(fixture)
    opts = X.ExportOptions(hash_source=False,
                           artefacts={"regions": True, "phases": False,
                                      "particles": False, "definitions": False,
                                      "xlsx": False})
    res = X.write_export(tables, tmp_path, options=opts)
    with open(Path(res.folder) / "regions.csv", newline="",
              encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    row = rows[0]
    assert float(row["frac_of_scan"]) == pytest.approx(
        tables.regions.rows[0]["frac_of_scan"], rel=0, abs=0)
    # 73/90 is not representable; the written text must round-trip exactly.
    assert float(row["frac_of_valid_eds"]) == N_REGION0_VALID / N_VALID


def test_formatting_a_cell(fixture):
    assert X._format_cell(None, ".") == ""
    assert X._format_cell(True, ".") == "True"
    assert X._format_cell(3, ".") == "3"
    assert X._format_cell(1 / 3, ".") == repr(1 / 3)
    assert X._format_cell(1 / 3, ",") == repr(1 / 3).replace(".", ",")
    assert X._format_cell(float("nan"), ".") == ""
    assert X._format_cell(X.NOT_MEASURED, ",") == X.NOT_MEASURED


# ---------------------------------------------------------------------------
# writing
# ---------------------------------------------------------------------------

def test_an_export_writes_every_always_on_artefact(tmp_path, fixture):
    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    names = {f.name for f in res.files}
    assert names == {"phases.csv", "regions.csv", "particles.csv",
                     "definitions.csv", "provenance.json", "summary.xlsx",
                     "phase_map.png", "region_map.png",
                     "phase_map_figure.png", "region_map_figure.png",
                     "caption.txt"}
    folder = Path(res.folder)
    assert folder.name.startswith("fixture_EDS_")


def test_the_opt_in_artefacts_are_off_by_default(tmp_path, fixture):
    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    names = {f.name for f in res.files}
    assert "labels.npz" not in names
    assert "pixels.csv" not in names


def test_an_existing_folder_is_never_overwritten(tmp_path, fixture):
    tables = _build(fixture)
    when = X.datetime(2026, 8, 27, 14, 30)
    opts = X.ExportOptions(hash_source=False, timestamp=when)
    first = X.write_export(tables, tmp_path, options=opts)
    second = X.write_export(tables, tmp_path, options=opts)
    assert Path(first.folder).name == "fixture_EDS_2026-08-27_1430"
    assert Path(second.folder).name == "fixture_EDS_2026-08-27_1430_2"
    # And the first one is untouched.
    assert (Path(first.folder) / "phases.csv").exists()


def test_on_existing_fail_raises_rather_than_suffixing(tmp_path, fixture):
    tables = _build(fixture)
    when = X.datetime(2026, 8, 27, 14, 30)
    X.write_export(tables, tmp_path,
                   options=X.ExportOptions(hash_source=False, timestamp=when))
    with pytest.raises(FileExistsError):
        X.write_export(tables, tmp_path,
                       options=X.ExportOptions(hash_source=False, timestamp=when,
                                               on_existing="fail"))


def test_provenance_json_is_written_and_is_valid_json(tmp_path, fixture):
    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    prov = json.loads((Path(res.folder) / "provenance.json").read_text(
        encoding="utf-8"))
    assert prov["format_version"] == X.FORMAT_VERSION
    assert prov["counts"]["n_px_total"] == N_PX


def test_the_workbook_opens_and_its_sheets_are_the_tables(tmp_path, fixture):
    from openpyxl import load_workbook

    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")
    # `warnings` is FIRST: a workbook opens on its first sheet, and the whole
    # purpose of a warning is to be met before the number it is about.
    assert wb.sheetnames == ["warnings", "phases", "regions", "particles",
                             "definitions", "provenance"]
    ws = wb["regions"]
    assert [c.value for c in ws[1]][:3] == [
        "region_id", "region_map_colour_hex", "n_px"]
    # Numbers go in as numbers so Excel applies the reader's own locale.
    assert isinstance(ws.cell(row=2, column=3).value, int)      # n_px


def test_labels_npz_round_trips_without_allow_pickle(tmp_path, fixture):
    """The method developer's highest-value artefact, loaded the safe way."""
    tables = _build(fixture)
    opts = X.ExportOptions(hash_source=False, artefacts={
        "phases": True, "regions": True, "particles": True,
        "definitions": True, "xlsx": False, "labels": True, "pixels": False})
    res = X.write_export(tables, tmp_path, options=opts)
    with np.load(Path(res.folder) / "labels.npz", allow_pickle=False) as z:
        assert z["phase_id"].shape == (N_ROWS, N_COLS)
        assert z["region_id"].shape == (N_ROWS, N_COLS)
        assert int(z["particle_id"].max()) == N_PARTICLES
        assert list(z["phase_lut_name"]) == ["Al.cif", "Si.cif", "Cu2Si.cif"]
        assert int(z["locked_mask"].sum()) == N_HAND


def test_the_particle_raster_and_the_particle_table_are_the_same_objects(
        tmp_path, fixture):
    tables = _build(fixture)
    opts = X.ExportOptions(hash_source=False, artefacts={
        "phases": False, "regions": False, "particles": True,
        "definitions": False, "xlsx": False, "labels": True, "pixels": False})
    res = X.write_export(tables, tmp_path, options=opts)
    with np.load(Path(res.folder) / "labels.npz", allow_pickle=False) as z:
        raster = z["particle_id"]
    for row in tables.particles.rows:
        assert int((raster == row["particle_id"]).sum()) == row["n_px"]


def test_pixels_csv_is_opt_in_and_has_one_row_per_pixel(tmp_path, fixture):
    tables = _build(fixture)
    opts = X.ExportOptions(hash_source=False, artefacts={
        "phases": False, "regions": False, "particles": False,
        "definitions": False, "xlsx": False, "labels": False, "pixels": True})
    res = X.write_export(tables, tmp_path, options=opts)
    written = {f.name: f for f in res.files}
    assert written["pixels.csv"].rows == N_PX
    with open(Path(res.folder) / "pixels.csv", newline="",
              encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == N_PX
    assert rows[0]["row"] == "0" and rows[0]["col"] == "0"
    assert {r["at_pct_Cu"] for r in rows} == {X.NOT_MEASURED}
    assert sum(1 for r in rows if r["hand_edited"] == "True") == N_HAND
    assert sum(1 for r in rows if r["has_eds"] == "False") == N_DEAD


def test_pixels_is_a_generator_and_is_not_materialised_by_build_tables(fixture):
    """A 485k-pixel map with twelve elements is ~100 MB of text. The table
    must not exist as a list of dicts before it is written."""
    tables = _build(fixture)
    assert callable(tables.pixel_rows)
    import types
    assert isinstance(tables.pixel_rows(), types.GeneratorType)


# ---------------------------------------------------------------------------
# definitions
# ---------------------------------------------------------------------------

def test_definitions_csv_is_written_even_when_there_are_none(tmp_path, fixture):
    tables = _build(fixture)
    assert len(tables.definitions) == 0
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    text = (Path(res.folder) / "definitions.csv").read_text(encoding="utf-8-sig")
    assert text.strip().startswith("kind,index,name")


def test_a_definition_is_exported_verbatim_with_what_it_claimed(fixture):
    state, at_maps = fixture
    state.region_defs = [{
        "name": "Si rich", "phase_key": "Si.cif",
        "elements": [{"element": "Si", "min_at_pct": 50.0, "max_at_pct": None}],
    }]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r = tables.definitions.rows[0]
    assert r["name"] == "Si rich"
    assert r["phase_key"] == "Si.cif"
    assert r["human_readable"] == "Si >= 50 at%"
    # Region 1 (60 at% Si, 13 px) and region 2 (55 at% Si, 4 px) both qualify.
    assert r["n_px_claimed"] == N_REGION1 + N_REGION2
    assert json.loads(r["clauses_json"])["elements"][0]["min_at_pct"] == 50.0


def test_a_region_a_definition_claims_says_definition_not_auto(fixture):
    state, at_maps = fixture
    state.region_defs = [{
        "name": "Si rich",
        "elements": [{"element": "Si", "min_at_pct": 58.0, "max_at_pct": None}],
    }]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r1 = _row(tables.regions, region_id=1)          # 60 at% Si — claimed
    r2 = _row(tables.regions, region_id=2)          # 55 at% Si — not claimed
    assert r1["assignment_source"] == "definition"
    assert r1["definition_name"] == "Si rich"
    assert r1["definition_match_frac"] == pytest.approx(1.0)
    assert r2["assignment_source"] == "auto"
    assert r2["definition_name"] == ""


def test_a_weak_overlap_is_recorded_but_not_claimed_as_the_source(fixture):
    """A window covering a third of a region is information; calling the
    region "definition" on the strength of it would be an overstatement."""
    state, at_maps = fixture
    # Claims the 3x3 block only (9 of region 1's 13 px = 0.69) — over the
    # half mark, so tighten it to the 2x2 piece instead: 4/13 = 0.31.
    state.region_defs = [{
        "name": "corner",
        "elements": [{"element": "Si", "min_at_pct": 59.0, "max_at_pct": 61.0}],
    }]
    # Make only the 2x2 piece match by pushing the 3x3 block off the window.
    flat = state.region_grid.ravel()
    block = np.zeros(N_PX, dtype=bool)
    block.reshape(N_ROWS, N_COLS)[1:4, 1:4] = True
    at_maps["Si"][block] = 70.0
    at_maps["Al"][block] = 30.0
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r1 = _row(tables.regions, region_id=1)
    assert r1["definition_match_frac"] == pytest.approx(4 / N_REGION1)
    assert r1["assignment_source"] == "auto"
    assert r1["definition_name"] == "corner"
    assert int(flat.max()) == 2                    # fixture unchanged otherwise


def test_hand_beats_definition_as_the_assignment_source(fixture):
    """A hand-painted pixel is the strongest statement in the file about
    where a boundary is."""
    state, at_maps = fixture
    state.region_defs = [{
        "name": "everything Al",
        "elements": [{"element": "Al", "min_at_pct": 50.0, "max_at_pct": None}],
    }]
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    r0 = _row(tables.regions, region_id=0)
    assert r0["definition_name"] == "everything Al"
    # `mixed` is the graded form of `hand` — 4 painted pixels of 83 — and it
    # still outranks a definition, for the same reason `hand` did: a painted
    # pixel is the strongest statement in the file about a boundary.
    assert r0["assignment_source"] == "mixed"


# ---------------------------------------------------------------------------
# a map with no regions at all
# ---------------------------------------------------------------------------

def test_a_pixel_mode_map_still_gets_particles_and_says_where_from(fixture):
    """Returning an empty particles.csv on every per-pixel map would be
    useless rather than honest. The basis is recorded because the two are not
    comparable."""
    state, at_maps = fixture
    state.region_grid = None
    state.region_phase = []
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert len(tables.regions) == 0
    assert len(tables.particles) > 0
    assert tables.provenance["particles"]["basis"] == "phase"
    assert any(w["code"] == "particles_from_phases" for w in tables.warnings)
    assert tables.particles.rows[0]["region_id"] is None
    assert tables.particles.rows[0]["phase_name"] in ("Al.cif", "Si.cif")


# ---------------------------------------------------------------------------
# preset override lands in the file
# ---------------------------------------------------------------------------

def test_an_overridden_refusal_is_recorded_in_the_provenance(fixture):
    """"A warning that only appears on screen is clicked away in half a
    second." """
    report = {"ok": False, "preset_name": "6xxx", "preset_hash": "abc123",
              "blockers": [{"code": "missing_element",
                            "message": "this scan never measured Cu"}],
              "warnings": []}
    tables = _build(fixture, preset_name="6xxx", compatibility=report)
    p = tables.provenance["preset"]
    assert p["name"] == "6xxx"
    assert p["override"] is True
    assert p["compatibility"]["blockers"][0]["code"] == "missing_element"


def test_no_preset_means_no_override_claim(fixture):
    p = _build(fixture).provenance["preset"]
    assert p["name"] is None
    assert p["override"] is False


# ---------------------------------------------------------------------------
# degenerate inputs must not produce plausible nonsense
# ---------------------------------------------------------------------------

def test_a_map_with_no_eds_at_all_reports_empty_fractions_not_zeros(fixture):
    state, _ = fixture
    dead = {"Al": np.zeros(N_PX), "Si": np.zeros(N_PX)}
    tables = X.build_tables(state, dead, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert tables.provenance["counts"]["n_px_valid_eds"] == 0
    assert any(w["code"] == "no_valid_eds" for w in tables.warnings)
    for r in tables.regions.rows:
        assert r["frac_of_valid_eds"] is None
        assert r["frac_of_scan"] is not None


def test_an_element_map_of_the_wrong_length_is_dropped_with_a_warning(fixture):
    state, at_maps = fixture
    at_maps["Fe"] = np.zeros(7)
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert any(w["code"] == "element_shape_mismatch" for w in tables.warnings)
    # Dropped from the measured set — so it reads `not measured`, which is
    # true: nothing usable was measured for it on this grid.
    assert "Fe" not in tables.provenance["elements"]["measured"]


def test_the_folder_name_keeps_the_first_dot_stem(tmp_path):
    when = X.datetime(2026, 8, 27, 9, 5)
    p = X.resolve_folder(tmp_path, "/data/SampleB_extrusion.h5oina", when)
    assert p.name == "SampleB_extrusion_EDS_2026-08-27_0905"
    assert X.resolve_folder(tmp_path, None, when).name == \
        "scan_EDS_2026-08-27_0905"


# ---------------------------------------------------------------------------
# plausibility: the export saying that its own headline number is wrong
#
# The case these pin is real. A group leader read a finished export and found
# 61 % of an aluminium extrusion classified as Fe-intermetallic by phases
# whose pixels carry a third of the iron those phases are made of - both
# numbers already in adjacent columns, and nothing subtracting them. The
# numbers quoted below are the measured ones from that scan (SampleB, 90x120,
# mode=cluster, scale_um=1.5), reproduced through the running backend.
# ---------------------------------------------------------------------------

def test_the_check_reproduces_the_two_rows_the_group_leader_challenged():
    """The acceptance case, in the units the file prints.

    ``sd_0302719`` measures 4.35 at% Fe (scored basis) against a nominal
    11.60, and ``Al6Fe`` 4.03 against 14.29. The map background for iron is
    2.86 at%, so iron is discriminating for both - the whole point being that
    aluminium, off by comparable ABSOLUTE amounts, is not.
    """
    bg = {"Al": 85.6605, "Fe": 2.8633, "Mn": 1.2569, "Si": 7.3406}
    measured_el = ["Al", "Cu", "Fe", "Mn", "Si", "Zn"]

    alfemnsi = X.composition_plausibility(
        {"Al": 80.2726, "Cu": 0.71765, "Fe": 4.26625, "Mn": 2.29741,
         "Si": 10.1669, "Zn": 0.2768},
        {"Al": 75.7714, "Fe": 11.6, "Mn": 2.68571, "Si": 9.94285},
        bg, measured_el, matrix_element="Al")
    assert alfemnsi["worst_element"] == "Fe"
    assert alfemnsi["worst_element_measured_at_pct"] == pytest.approx(4.353, abs=1e-3)
    assert alfemnsi["worst_element_nominal_at_pct"] == pytest.approx(11.6)
    assert alfemnsi["worst_element_ratio"] == pytest.approx(2.665, abs=1e-3)
    assert alfemnsi["worst_element_direction"] == "deficient"
    assert alfemnsi["composition_implausible"] is True

    al6fe = X.composition_plausibility(
        {"Al": 82.4267, "Cu": 0.68424, "Fe": 3.95087, "Mn": 2.12264,
         "Si": 8.63627, "Zn": 0.2718},
        {"Al": 85.7142, "Fe": 14.2857},
        bg, measured_el, matrix_element="Al")
    assert al6fe["worst_element"] == "Fe"
    assert al6fe["worst_element_measured_at_pct"] == pytest.approx(4.028, abs=1e-3)
    assert al6fe["worst_element_ratio"] == pytest.approx(3.547, abs=1e-3)
    assert al6fe["composition_implausible"] is True


def test_the_good_elements_of_the_same_bad_row_are_not_flagged():
    """The criterion is per element, not a verdict on the row.

    ``sd_0302719`` fails at 2.665x on iron. Its manganese - also
    discriminating, at 2.69 at% nominal over a 1.26 at% background - sits at
    1.146x, which is the largest ratio a good call produces anywhere on that
    scan and is the measurement the 2.0 threshold is placed above.
    """
    disc = X.discriminating_elements(
        {"Al": 75.7714, "Fe": 11.6, "Mn": 2.68571, "Si": 9.94285},
        {"Al": 85.6605, "Fe": 2.8633, "Mn": 1.2569, "Si": 7.3406},
        matrix_element="Al")
    # Si is IN the phase at 9.94 at% and is still not discriminating: the map
    # already carries 7.34 at% of it everywhere (1.35x, under 2.0).
    assert disc == ["Fe", "Mn"]

    meas = X.to_scored_basis(
        {"Al": 80.2726, "Cu": 0.71765, "Fe": 4.26625, "Mn": 2.29741,
         "Si": 10.1669, "Zn": 0.2768})
    mn = max(2.68571, meas["Mn"]) / min(2.68571, meas["Mn"])
    assert mn == pytest.approx(1.146, abs=1e-3)
    assert mn < X._IMPLAUSIBLE_FACTOR


def test_the_matrix_element_never_discriminates_however_it_is_ruled_out():
    """Two independent reasons, because either alone can be wrong.

    On an Al-matrix scan aluminium falls out on the background test on its own
    (100 / 85.66 = 1.17x, under 2.0) - so the check does not depend on
    ``infer_matrix_element`` getting the solvent right. And on a map whose
    matrix is NOT its most abundant background the explicit rule catches it,
    which is why both are here.
    """
    # background test alone, matrix not declared
    assert X.discriminating_elements(
        {"Al": 100.0}, {"Al": 85.6605}, matrix_element=None) == []
    # explicit rule alone: a matrix that would pass the background test
    assert X.discriminating_elements(
        {"Al": 90.0, "Fe": 10.0}, {"Al": 30.0, "Fe": 1.0},
        matrix_element="Al") == ["Fe"]


def test_a_trace_element_of_the_phase_cannot_carry_the_verdict():
    """An element a phase needs at a few tenths of a percent is below what
    this semi-quantitative pipeline can resolve, so a factor on it is noise
    wearing a decimal point. The smallest discriminating nominal anywhere on
    the real scan is manganese at 2.69 at%, well clear of the 1.0 floor.

    Without the floor this phase would read 10x implausible off a nominal
    0.5 at% against a measured nothing, and be the loudest row in the file.
    """
    nominal = {"Al": 99.5, "Fe": 0.5}
    bg = {"Al": 85.0, "Fe": 0.01}
    assert X.discriminating_elements(nominal, bg, matrix_element="Al") == []
    cells = X.composition_plausibility(
        {"Al": 100.0, "Fe": 0.0}, nominal, bg, ["Al", "Fe"],
        matrix_element="Al")
    assert cells["plausibility_checked"] is False
    assert cells["composition_implausible"] is None
    # ...and one at% more of the same element IS checkable, so the floor is
    # the only thing standing between the two answers.
    assert X.discriminating_elements({"Al": 98.5, "Fe": 1.5}, bg,
                                     matrix_element="Al") == ["Fe"]


def test_the_background_is_read_in_at_pct_and_not_in_shares():
    """The units regression, and the reason that function exists.

    ``chemistry_score.background_levels`` returns the median renormalised
    SHARE (0..1), not a percentage. The first cut of this check compared a
    nominal 11.60 at% against 0.0286 and printed "map background 0.03 at%"
    for an iron level the export's own ``background_at_pct_Fe`` column gives
    as 2.8633. At that scale every element clears the discriminating test and
    the whole criterion is silently off.
    """
    shares = {"Al": 0.856605, "Fe": 0.028633}
    assert X.background_at_pct(shares)["Fe"] == pytest.approx(2.8633)
    # Aluminium is discriminating against the SHARE and is not against the
    # at%. That difference is the bug, in one assertion.
    assert X.discriminating_elements({"Al": 100.0}, shares) == ["Al"]
    assert X.discriminating_elements({"Al": 100.0},
                                     X.background_at_pct(shares)) == []


def test_the_background_the_check_divides_by_is_the_column_the_file_prints(fixture):
    """One background, so a warning cannot contradict the row beside it."""
    _state, at_maps = fixture
    at_pct = X.background_at_pct(X._background_levels(at_maps))
    tables = _build(fixture)
    row = _row(tables.phases, phase_name="Si.cif")
    for el, v in at_pct.items():
        assert row[f"background_at_pct_{el}"] == pytest.approx(v)


def test_a_phase_whose_elements_all_sit_at_background_is_not_checked(fixture):
    """"Cannot be checked" is not "checked and fine".

    ``Al.cif`` on an aluminium matrix has no element that stands above the map
    background, so chemistry has nothing to say about the assignment. The
    verdict is ``None``, never ``False``, and ``plausibility_note`` says why -
    the same distinction NOT_MEASURED draws between a measured zero and never
    having looked.
    """
    tables = _build(fixture)
    row = _row(tables.phases, phase_name="Al.cif")
    assert row["plausibility_checked"] is False
    assert row["composition_implausible"] is None
    assert row["review_flag"] is False
    assert "background" in row["plausibility_note"]


def test_a_good_match_does_not_trip(fixture):
    """The fixture's silicon regions are a correct call and must stay quiet.

    17 px of Si.cif at a mean 58.82 at% Si against a nominal 100 - a real
    interaction-volume dilution, 1.70x, under the threshold.
    """
    tables = _build(fixture)
    row = _row(tables.phases, phase_name="Si.cif")
    assert row["plausibility_checked"] is True
    assert row["worst_element"] == "Si"
    assert row["worst_element_measured_at_pct"] == pytest.approx(1000 / 17)
    assert row["worst_element_ratio"] == pytest.approx(100 * 17 / 1000)
    assert row["composition_implausible"] is False
    assert row["review_flag"] is False
    assert not any(w["code"].startswith("composition_implausible")
                   for w in tables.warnings)


def test_an_exact_composition_gives_a_ratio_of_exactly_one():
    cells = X.composition_plausibility(
        {"Al": 85.7143, "Fe": 14.2857}, {"Al": 85.7143, "Fe": 14.2857},
        {"Al": 85.0, "Fe": 2.0}, ["Al", "Fe"])
    assert cells["worst_element_ratio"] == pytest.approx(1.0)
    assert cells["worst_element_direction"] == "exact"
    assert cells["composition_implausible"] is False


def test_the_criterion_is_symmetric_in_direction():
    """A phase assigned to pixels carrying far MORE of its discriminating
    element than it contains is the same kind of wrong, and a one-sided
    ``nominal / measured`` would report 0.33 and call it fine."""
    lots = X.composition_plausibility(
        {"Al": 70.0, "Fe": 30.0}, {"Al": 90.0, "Fe": 10.0},
        {"Al": 85.0, "Fe": 2.0}, ["Al", "Fe"])
    assert lots["worst_element"] == "Fe"
    assert lots["worst_element_ratio"] == pytest.approx(3.0)
    assert lots["worst_element_direction"] == "excess"
    assert lots["composition_implausible"] is True


def test_exactly_at_the_threshold_counts_as_implausible():
    """The comparison is ``>=``. A boundary that silently excludes itself is
    the kind of off-by-one nobody ever sees in a report."""
    cells = X.composition_plausibility(
        {"Al": 95.0, "Fe": 5.0}, {"Al": 90.0, "Fe": 10.0},
        {"Al": 85.0, "Fe": 2.0}, ["Al", "Fe"])
    assert cells["worst_element_ratio"] == pytest.approx(X._IMPLAUSIBLE_FACTOR)
    assert cells["composition_implausible"] is True


def test_a_measured_zero_gives_a_large_finite_ratio_not_an_infinity():
    """0.05 at% is the smallest value the region inspector shows; below it
    there is no measurement to divide by. The answer stays a number a
    spreadsheet can sort."""
    cells = X.composition_plausibility(
        {"Al": 100.0, "Fe": 0.0}, {"Al": 62.5, "Fe": 37.5},
        {"Al": 85.0, "Fe": 2.0}, ["Al", "Fe"])
    ratio = cells["worst_element_ratio"]
    assert math.isfinite(ratio)
    assert ratio == pytest.approx(37.5 / X._MEASURED_FLOOR_AT_PCT)


def test_an_element_the_phase_needs_but_nobody_mapped_is_renormalised_away():
    """Both sides of the comparison have to share a denominator.

    A phase containing chromium this scan never mapped would otherwise claim
    a share of a total the measurement cannot account for, and every other
    element of that phase would read systematically low against it. Same
    lesson as ``to_scored_basis``, mirrored.
    """
    nominal = X.nominal_on_measured_basis(
        {"Al": 50.0, "Fe": 25.0, "Cr": 25.0}, ["Al", "Fe"])
    assert nominal == pytest.approx({"Al": 200 / 3, "Fe": 100 / 3})
    # identity when the scan measured everything the phase names
    assert X.nominal_on_measured_basis({"Al": 60.0, "Fe": 40.0},
                                       ["Al", "Fe", "Si"]) == \
        pytest.approx({"Al": 60.0, "Fe": 40.0})
    # and carbon/oxygen are dropped like everywhere else
    assert "O" not in X.nominal_on_measured_basis(
        {"Al": 40.0, "O": 60.0}, ["Al", "O"], ignore={"O"})


def test_a_row_with_no_phase_carries_no_verdict(fixture):
    tables = _build(fixture)
    row = _row(tables.phases, phase_name="unclassified")
    assert row["plausibility_checked"] is False
    assert row["composition_implausible"] is None


# ---------------------------------------------------------------------------
# a fixture that is genuinely wrong, so the warning path is exercised
# ---------------------------------------------------------------------------

@pytest.fixture
def bad_fixture(fixture):
    """The same map with the silicon regions diluted to 20 at% Si.

    Si.cif then needs 100 at% silicon on pixels reading 20, a factor of 5 -
    the same shape of failure as the real scan's 2.7x iron, and 17 % of the
    raster, well over the 5 % naming floor.
    """
    state, at_maps = fixture
    flat_region = np.asarray(state.region_grid).ravel()
    hit = flat_region > 0
    at_maps["Si"][hit] = 20.0
    at_maps["Al"][hit] = 80.0
    return state, at_maps


def test_a_big_implausible_row_is_named_in_the_warnings(bad_fixture):
    tables = _build(bad_fixture)
    row = _row(tables.phases, phase_name="Si.cif")
    assert row["worst_element_ratio"] == pytest.approx(5.0)
    assert row["composition_implausible"] is True
    assert row["review_flag"] is True
    assert row["review_reasons"] == "implausible_composition"

    named = [w for w in tables.warnings if w["code"] == "composition_implausible"]
    assert len(named) == 1
    msg = named[0]["message"]
    # The message has to carry every number a reader needs to check it.
    for piece in ("Si.cif", "17.00 %", "20.00 at% Si", "100.00 at%", "5.00x",
                  "deficient", "Nothing was changed"):
        assert piece in msg, msg


def test_a_tiny_implausible_row_is_flagged_but_not_named(bad_fixture):
    """Weighting by coverage. A 4 px row is a curiosity; it must not shout,
    and it must not vanish either - it is flagged in the table and counted in
    the summary."""
    state, at_maps = bad_fixture
    # Shrink the silicon to region 2 alone: 4 px = 4 % of the raster, under
    # the 5 % naming floor.
    flat_region = np.asarray(state.region_grid).ravel()
    at_maps["Si"][flat_region == 1] = 5.0
    at_maps["Al"][flat_region == 1] = 95.0
    state.region_phase = [0, 0, 1]
    phase = np.asarray(state.phase_grid).ravel().copy()
    phase[flat_region == 1] = 0                 # region 1 is matrix again
    state.phase_grid = phase.reshape(N_ROWS, N_COLS).astype(np.int32)

    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    row = _row(tables.phases, phase_name="Si.cif")
    assert row["frac_of_scan"] < X._WARN_MIN_FRAC_OF_SCAN
    assert row["composition_implausible"] is True

    assert not [w for w in tables.warnings
                if w["code"] == "composition_implausible"]
    summary = [w for w in tables.warnings
               if w["code"] == "composition_implausible_summary"]
    assert len(summary) == 1
    assert "flagged in the tables only" in summary[0]["message"]
    assert tables.provenance["plausibility"]["n_rows_implausible"] == 1
    assert tables.provenance["plausibility"]["n_rows_named_in_warnings"] == 0


def test_the_warning_is_raised_from_the_phases_table_only(bad_fixture):
    """A region and the phase it carries are largely the same pixels.

    Warning from both tables would state one finding twice and double the
    fraction of the map it claims to describe. The regions still carry the
    per-row columns, which is where the detail lives.
    """
    tables = _build(bad_fixture)
    assert len([w for w in tables.warnings
                if w["code"] == "composition_implausible"]) == 1
    flagged = [r for r in tables.regions.rows
               if r["composition_implausible"] is True]
    assert len(flagged) == 2          # both silicon regions carry the flag
    total = tables.provenance["plausibility"]["frac_of_scan_implausible"]
    assert total == pytest.approx(N_REGION1 / N_PX + N_REGION2 / N_PX)


def test_the_provenance_carries_the_criterion_and_says_it_changes_nothing(
        bad_fixture):
    tables = _build(bad_fixture)
    p = tables.provenance["plausibility"]
    assert p["thresholds"]["implausible_factor"] == X._IMPLAUSIBLE_FACTOR
    assert p["thresholds"]["discriminating_factor_over_background"] == \
        X._DISCRIMINATING_FACTOR
    assert p["thresholds"]["warn_min_frac_of_scan"] == X._WARN_MIN_FRAC_OF_SCAN
    assert "report on the result" in p["changes_nothing"]
    assert set(p["review_reasons"]) == set(X.REVIEW_REASONS)
    assert p["rows"][0]["element"] == "Si"


def test_the_check_is_a_report_and_never_moves_an_assignment(bad_fixture,
                                                             monkeypatch):
    """Mutation of the criterion must not reach a single other cell.

    Rebuilt with the threshold moved from 2.0 to 99, which turns every verdict
    off. If the check leaked into the classification - a gate, a re-score, a
    tie-break - some column outside the plausibility block would move with it.
    """
    before = _build(bad_fixture)
    monkeypatch.setattr(X, "_IMPLAUSIBLE_FACTOR", 99.0)
    after = _build(bad_fixture)

    plaus = set(X.plausibility_columns())
    for t_before, t_after in ((before.phases, after.phases),
                              (before.regions, after.regions),
                              (before.particles, after.particles)):
        assert t_before.columns == t_after.columns
        assert len(t_before.rows) == len(t_after.rows)
        for a, b in zip(t_before.rows, t_after.rows):
            for col in t_before.columns:
                if col in plaus:
                    continue
                assert a.get(col) == b.get(col) or (
                    a.get(col) != a.get(col) and b.get(col) != b.get(col)), col

    # ...and the verdict itself did move, so the mutation was real.
    assert _row(before.phases, phase_name="Si.cif")["composition_implausible"]
    assert not _row(after.phases, phase_name="Si.cif")["composition_implausible"]
    assert not [w for w in after.warnings
                if w["code"].startswith("composition_implausible")]


def test_the_verdict_sits_beside_the_phase_it_judges_in_all_three_tables(
        fixture):
    """``review_flag`` at column BW of 76 is a flag nobody scrolls to.

    That is what a QA tester measured, in the one table she filters row by
    row on a weekly batch. The block now sits immediately behind the row's
    phase identity — and the block stays CONTIGUOUS, so
    ``plausibility_columns()`` still describes something a reader can select
    as one range.

    The whole block is asserted to be contiguous and near the front rather
    than pinned to an exact index: an index would break on the next legitimate
    addition to the identity columns, and "near the front" is the property
    the tester actually complained about.
    """
    tables = _build(fixture)
    block = X.plausibility_columns()
    for table in (tables.phases, tables.regions, tables.particles):
        assert set(block) <= set(table.columns), table.name
        at = table.columns.index(block[0])
        # contiguous, in order
        assert table.columns[at:at + len(block)] == block, table.name
        # and where a person actually looks. The regions table carries the
        # widest identity block (five phase_final_* columns plus the counts
        # and areas that precede them), which is why the bound is not tighter.
        assert at <= 20, (table.name, at)
        # ...which is a long way in front of where it used to be.
        assert at < len(table.columns) - len(block)


def test_moving_the_block_cost_a_format_version(fixture):
    """The export's own compatibility rule, applied rather than dodged.

    "Additions only at the end" says an APPEND is free and anything else
    breaks loudly. Moving the block is not an append, so the pin moves: a
    position-pinned reader of version 3 now fails on the version it pins
    instead of silently reading ``worst_element`` as ``mean_at_pct_Al``.
    Every column keeps its name, so a reader that selects by HEADER —
    ``csv.DictReader``, every reader in this repo — is unaffected.
    """
    assert X.FORMAT_VERSION == 4
    tables = _build(fixture)
    assert tables.provenance["format_version"] == 4
    # names unchanged, which is what makes a by-header reader safe
    assert X.plausibility_columns() == [
        "plausibility_checked", "plausibility_note",
        "worst_element", "worst_element_measured_at_pct",
        "worst_element_nominal_at_pct", "worst_element_ratio",
        "worst_element_direction", "composition_implausible",
        "review_flag", "review_reasons",
    ]


def test_review_flag_ors_the_reasons_the_row_can_actually_answer(fixture):
    """A missing column must not read as "fine".

    ``regions`` has a margin and an assignment source; ``phases`` has neither,
    so it contributes only the composition reason.
    """
    assert X._review_reasons({"composition_implausible": True}) == \
        ["implausible_composition"]
    # 0.0048 is the real margin of the Al6Fe region, 27.64 % of the SampleB
    # map - a coin toss wearing a phase name.
    assert X._review_reasons({"margin_score": 0.0048}) == \
        ["undecided_between_phases"]
    assert X._review_reasons({"margin_score": 0.0492}) == []
    assert X._review_reasons({"assignment_source": "untracked"}) == \
        ["untracked_assignment"]
    assert X._review_reasons({"assignment_source": "hand"}) == []
    assert X._review_reasons({}) == []
    both = X._review_reasons({"composition_implausible": True,
                              "margin_score": 0.001,
                              "assignment_source": "untracked"})
    assert both == ["implausible_composition", "undecided_between_phases",
                    "untracked_assignment"]


def test_the_plausibility_columns_survive_a_write(tmp_path, bad_fixture):
    tables = _build(bad_fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    with open(Path(res.folder) / "phases.csv", newline="",
              encoding="utf-8-sig") as fh:
        rows = {r["phase_name"]: r for r in csv.DictReader(fh)}
    assert rows["Si.cif"]["composition_implausible"] == "True"
    assert float(rows["Si.cif"]["worst_element_ratio"]) == pytest.approx(5.0)
    assert rows["Si.cif"]["review_reasons"] == "implausible_composition"
    # "cannot be checked" writes an EMPTY cell, not False.
    assert rows["Al.cif"]["composition_implausible"] == ""


# ---------------------------------------------------------------------------
# the fraction that passed NOTHING
# ---------------------------------------------------------------------------
#
# WHY THIS BLOCK EXISTS. The export reported both halves of the answer and
# never added them up. On the real 90x120 SampleB scan the warning read "2 of
# 2 checkable phase rows are composition-implausible ... covering 34.69 % of
# the scan", and the tester who filed this said a hurried reader takes that
# as "a third is dodgy, two-thirds are fine". The two-thirds are 65.31 % of a
# map chemistry cannot check at all, and the corroborated fraction is 0.00 %.
# Nothing in that map was corroborated. These tests hold the three numbers
# together, and hold the summary sentence to stating all three.

def test_the_three_coverages_partition_the_scan(fixture):
    """implausible + not checkable + corroborated == the whole map.

    This is the property that makes stating all three honest rather than
    merely fuller: ``_phase_table`` emits one row per phase that owns pixels
    PLUS a real ``unclassified`` row, and those masks cover every pixel
    exactly once. If the sum ever stops closing, one of the three is being
    counted against a different set than the other two — which is the
    original defect one level up.
    """
    tables = _build(fixture)
    p = tables.provenance["plausibility"]
    total = (p["frac_of_scan_implausible"]
             + p["frac_of_scan_not_checkable"]
             + p["frac_of_scan_corroborated"])
    assert total == pytest.approx(1.0)
    assert p["frac_of_scan_accounted"] == pytest.approx(1.0)
    assert p["n_rows_without_frac"] == 0
    assert (p["n_rows_implausible"] + p["n_rows_corroborated"]
            + p["n_rows_not_checkable"]) == len(tables.phases.rows)
    assert p["n_rows_checkable"] == (p["n_rows_implausible"]
                                     + p["n_rows_corroborated"])


def test_the_not_checkable_fraction_is_the_number_that_was_missing(fixture):
    """The fixture, counted by hand.

    Three phase rows own pixels here, and the counts are the fixture's own
    constants rather than round numbers:

      Al.cif          69 px  = 83 matrix px, less the 10-pixel dead strip and
                             the 4 hand-painted ones, both of which fall in
                             the matrix and are unclassified. NOT CHECKABLE:
                             no element of pure aluminium stands far enough
                             above an aluminium background to testify.
      Si.cif          17 px  = the two silicon regions. Checkable, and correct.
      unclassified    14 px  = the dead strip plus the hand-painted pixels.
                             No phase, so nothing to check it against.

    So 69 + 14 = 83 px of the 100 px raster cannot be checked at all and 17 %
    is corroborated — and the three still close on the whole map, which the
    partition test above asserts separately.
    """
    tables = _build(fixture)
    p = tables.provenance["plausibility"]
    assert p["frac_of_scan_not_checkable"] == pytest.approx(0.83)
    assert p["frac_of_scan_corroborated"] == pytest.approx(0.17)
    assert p["frac_of_scan_implausible"] == pytest.approx(0.0)
    assert p["n_rows_not_checkable"] == 2          # Al.cif and unclassified
    assert p["n_rows_corroborated"] == 1           # Si.cif

    # The counts behind those fractions, read off the rows themselves so a
    # fixture change cannot leave the arithmetic above quietly meaningless.
    assert _row(tables.phases, phase_name="Al.cif")["n_px"] == 69
    assert _row(tables.phases, phase_name="Si.cif")["n_px"] == 17
    assert _row(tables.phases, phase_name="unclassified")["n_px"] == 14


def test_the_summary_warning_states_all_three_numbers(bad_fixture):
    """The sentence the tester asked for, with the arithmetic behind it.

    In ``bad_fixture`` the silicon regions are diluted to 20 at% Si, so
    Si.cif is the only checkable row and it fails: 17 % implausible, 83 % not
    checkable, and NOTHING corroborated. All three have to be in the one
    sentence, because the reader's error is made by subtraction.
    """
    tables = _build(bad_fixture)
    summary = [w for w in tables.warnings
               if w["code"] == "composition_implausible_summary"]
    assert len(summary) == 1
    msg = summary[0]["message"]
    for piece in ("17.00 % implausible", "83.00 % not checkable",
                  "0.00 % corroborated"):
        assert piece in msg, msg
    # ...and it says what "not checkable" IS, because that is the word the
    # old wording let a reader silently translate into "fine".
    assert "ABSENCE OF" in msg and "NOT A PASS" in msg
    # ...and the manual's conclusion is now in the folder, where the numbers
    # are, rather than only in a manual the reader may never open.
    assert "Nothing in this map was corroborated" in msg

    p = tables.provenance["plausibility"]
    assert p["frac_of_scan_corroborated"] == pytest.approx(0.0)
    assert p["frac_of_scan_not_checkable"] == pytest.approx(0.83)


def test_the_everything_failed_sentence_only_fires_when_everything_failed(
        fixture):
    """It is a claim about the whole map and must not be made loosely.

    Here one checkable row fails and another passes, so "every checkable row
    failed" is FALSE and that sentence must be absent — while the three
    fractions stay unconditional.
    """
    state, at_maps = fixture
    flat_region = np.asarray(state.region_grid).ravel()
    # Region 1 goes to Cu2Si.cif, which needs 33.33 at% Si on pixels that
    # will read 5 — a discriminating element, badly off. Region 2 keeps its
    # correct Si.cif. So: one implausible row and one corroborated one.
    at_maps["Si"][flat_region == 1] = 5.0
    at_maps["Al"][flat_region == 1] = 95.0
    state.region_phase = [0, 2, 1]
    phase = np.asarray(state.phase_grid).ravel().copy()
    phase[flat_region == 1] = 2
    state.phase_grid = phase.reshape(N_ROWS, N_COLS).astype(np.int32)

    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    p = tables.provenance["plausibility"]
    assert p["n_rows_implausible"] >= 1
    assert p["n_rows_corroborated"] >= 1
    assert p["n_rows_implausible"] < p["n_rows_checkable"]

    summary = [w for w in tables.warnings
               if w["code"] == "composition_implausible_summary"]
    assert len(summary) == 1
    msg = summary[0]["message"]
    assert "Nothing in this map was corroborated" not in msg
    for piece in ("% implausible", "% not checkable", "% corroborated"):
        assert piece in msg, msg


def test_a_row_whose_fraction_is_missing_is_counted_but_carries_no_area():
    """A row that cannot contribute area must not silently contribute 0 %.

    The three fractions are sold as a partition of the map. A row with no
    ``frac_of_scan`` breaks that, and the file has to say so rather than let
    the sum quietly come out short.
    """
    cov = X.plausibility_coverage([
        {"plausibility_checked": True, "composition_implausible": True,
         "frac_of_scan": 0.4},
        {"plausibility_checked": True, "composition_implausible": False,
         "frac_of_scan": 0.4},
        {"plausibility_checked": False, "composition_implausible": None,
         "frac_of_scan": None},
    ])
    assert cov["n_rows_total"] == 3
    assert cov["n_rows_without_frac"] == 1
    assert cov["frac_of_scan_accounted"] == pytest.approx(0.8)
    assert cov["frac_of_scan_not_checkable"] == pytest.approx(0.0)
    assert cov["n_rows_not_checkable"] == 1


def test_an_unchecked_row_is_never_counted_as_corroborated():
    """The distinction the whole complaint is about, at the unit level.

    ``composition_implausible`` is None on an unchecked row, and
    ``None is not True`` would land it in the corroborated bucket if the
    branch tested the verdict before testing whether anything was checked.
    """
    cov = X.plausibility_coverage([
        {"plausibility_checked": False, "composition_implausible": None,
         "frac_of_scan": 0.9},
        {"plausibility_checked": True, "composition_implausible": False,
         "frac_of_scan": 0.1},
    ])
    assert cov["frac_of_scan_corroborated"] == pytest.approx(0.1)
    assert cov["frac_of_scan_not_checkable"] == pytest.approx(0.9)
    assert cov["n_rows_corroborated"] == 1


def test_mutating_the_coverage_arithmetic_is_caught(bad_fixture, monkeypatch):
    """Mutation test for the number the testers singled out.

    A coverage function that folded the not-checkable rows into the
    corroborated ones — the exact mistake the old file made in prose — has to
    fail something. Injected here as a deliberately wrong
    ``plausibility_coverage``; both the provenance number AND the warning
    sentence must move, which is what proves the sentence is computed from
    the arithmetic rather than written beside it.
    """
    good = _build(bad_fixture)
    assert good.provenance["plausibility"]["frac_of_scan_corroborated"] == \
        pytest.approx(0.0)

    real = X.plausibility_coverage

    def _wrong(rows):
        out = dict(real(rows))
        out["frac_of_scan_corroborated"] += out["frac_of_scan_not_checkable"]
        out["frac_of_scan_not_checkable"] = 0.0
        return out

    monkeypatch.setattr(X, "plausibility_coverage", _wrong)
    bad = _build(bad_fixture)
    assert bad.provenance["plausibility"]["frac_of_scan_corroborated"] == \
        pytest.approx(0.83)
    msg = [w for w in bad.warnings
           if w["code"] == "composition_implausible_summary"][0]["message"]
    assert "83.00 % corroborated" in msg
    assert "0.00 % not checkable" in msg


def test_particles_carry_the_review_flag_too(bad_fixture):
    """The one table filtered row by row on a weekly batch.

    It was the only table without a ``review_flag``, and the QA tester's
    answer to "which table do you filter" was this one.
    """
    tables = _build(bad_fixture)
    assert "review_flag" in tables.particles.columns
    assert "review_reasons" in tables.particles.columns
    si = [r for r in tables.particles.rows if r["phase_name"] == "Si.cif"]
    assert si, "the fixture should have silicon particles"
    for r in si:
        assert r["composition_implausible"] is True
        assert r["review_flag"] is True
        assert "implausible_composition" in r["review_reasons"]
    # ...and the warnings still come from phases.csv alone: one wrong phase
    # must not shout once per connected component of it.
    assert len([w for w in tables.warnings
                if w["code"] == "composition_implausible"]) == 1


def test_the_particle_flag_is_the_particles_own_verdict(fixture):
    """Not the region's — the defect this module was already repaired for.

    Every particle here sits on correctly-assigned pixels, so nothing is
    flagged; and the matrix particles read ``plausibility_checked = False``,
    which is an absence of evidence and not a pass.
    """
    tables = _build(fixture)
    assert not [r for r in tables.particles.rows if r["review_flag"]]
    al = [r for r in tables.particles.rows if r["phase_name"] == "Al.cif"]
    assert al
    for r in al:
        assert r["plausibility_checked"] is False
        assert r["composition_implausible"] is None


def test_the_review_columns_survive_a_write_of_particles(tmp_path,
                                                         bad_fixture):
    tables = _build(bad_fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    with open(Path(res.folder) / "particles.csv", newline="",
              encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert rows
    si = [r for r in rows if r["phase_name"] == "Si.cif"]
    assert si and all(r["review_flag"] == "True" for r in si)


# ---------------------------------------------------------------------------
# provenance.json opens with the naive call
# ---------------------------------------------------------------------------

def test_provenance_json_loads_without_an_encoding_argument(tmp_path,
                                                            fixture):
    """The acceptance test, written literally the way the tester wrote it.

    ``json.load(open(p))`` uses the LOCALE default, which on Windows is
    cp1252. The file used to be UTF-8 carrying subscripted formulas out of
    the CIF library, so this exact call raised UnicodeDecodeError on the
    machine the tester was on. It is pure ASCII now.
    """
    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    p = Path(res.folder) / "provenance.json"

    loaded = json.load(open(p))            # noqa: SIM115 — that IS the test
    assert loaded["format_version"] == X.FORMAT_VERSION
    assert loaded["encoding"]["provenance_json"] == "ascii"

    # ...and it really is ASCII on disk, which is what makes the line above
    # work under every locale rather than only under this one.
    p.read_bytes().decode("ascii")


def test_a_subscripted_formula_round_trips_through_the_escapes(tmp_path,
                                                               fixture):
    """ASCII-escaping must not transliterate anything.

    ``\\u2081`` is valid JSON that every parser decodes back to the identical
    character, so the file is byte-for-byte ASCII AND character-for-character
    unchanged. Substituting ``formula_ascii`` would have been a different,
    LOSSY answer — that column exists, but beside the real one.
    """
    state, at_maps = fixture
    state.phase_entries[0].formula = "Mg₁₇Al₁₂"
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    p = Path(res.folder) / "provenance.json"

    raw = p.read_bytes()
    assert b"Mg\\u2081\\u2087Al\\u2081\\u2082" in raw     # escaped on disk...
    raw.decode("ascii")                                  # ...so it is ASCII

    loaded = json.load(open(p))
    entry = loaded["phases_participating"][0]
    assert entry["formula"] == "Mg₁₇Al₁₂"   # identical back
    assert entry["formula_ascii"] == "Mg17Al12"


# ---------------------------------------------------------------------------
# the warnings are readable in Excel
# ---------------------------------------------------------------------------

def test_every_warning_gets_its_own_wrapped_row(tmp_path, bad_fixture):
    """It was a 2074-character JSON blob in one unwrapped cell — B300 of a
    302-row sheet. The tester: the most important content in the export, and
    the worst-presented thing in it."""
    from openpyxl import load_workbook

    tables = _build(bad_fixture)
    assert tables.warnings, "the bad fixture should warn"
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")

    ws = wb["warnings"]
    assert [c.value for c in ws[1]] == ["#", "code", "message"]
    rows = [(r[1].value, r[2].value) for r in ws.iter_rows(min_row=2)]
    assert len(rows) == len(tables.warnings)
    for (code, message), w in zip(rows, tables.warnings):
        # Code and message in SEPARATE columns: one is what a reader filters
        # on, the other is what they read, and one cell holding both is
        # neither.
        assert code == w["code"]
        assert message == w["message"]
    # Wrapped, or a four-sentence message renders as one clipped line with an
    # invisible remainder — which is the whole complaint, one cell smaller.
    assert ws.cell(row=2, column=3).alignment.wrap_text is True
    assert ws.column_dimensions["C"].width >= 60


def test_the_provenance_sheet_points_at_the_warnings_sheet(tmp_path,
                                                           bad_fixture):
    """The blob is REPLACED, not duplicated. Leaving it in place beside the
    new sheet would keep the unreadable copy and add a readable one."""
    from openpyxl import load_workbook

    tables = _build(bad_fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")

    cells = {r[0].value: r[1].value
             for r in wb["provenance"].iter_rows(min_row=2)}
    assert "warnings" in cells
    assert "warnings' sheet" in str(cells["warnings"])
    # The long prose of an actual warning is NOT on the provenance sheet.
    long_message = max((w["message"] for w in tables.warnings), key=len)
    assert all(long_message not in str(v) for v in cells.values())

    # ...and provenance.json still carries the full list.
    loaded = json.load(open(Path(res.folder) / "provenance.json"))
    assert [w["code"] for w in loaded["warnings"]] == \
        [w["code"] for w in tables.warnings]


def test_an_export_with_nothing_to_warn_about_says_so_on_the_sheet(tmp_path,
                                                                   fixture):
    """A blank sheet reads as a broken export."""
    from openpyxl import load_workbook

    tables = _build(fixture)
    tables.warnings.clear()
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")
    ws = wb["warnings"]
    assert ws["B2"].value == "(none)"
    assert "no warnings" in str(ws["C2"].value)


def test_the_workbook_carries_warnings_raised_after_the_provenance_snapshot(
        tmp_path, fixture):
    """``_provenance`` snapshots the bag with ``list(warnings)``; the
    writer's bag is a superset. The sheet reads the LIVE bag, so a
    ``map_png_failed`` raised during writing still reaches the reader."""
    from openpyxl import load_workbook

    tables = _build(fixture)
    tables.warnings.append({
        "code": "raised_during_write",
        "message": "this was appended after the provenance snapshot"})
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")
    codes = [r[1].value for r in wb["warnings"].iter_rows(min_row=2)]
    assert "raised_during_write" in codes
    # ...and it is NOT in provenance.json, which was snapshotted earlier.
    loaded = json.load(open(Path(res.folder) / "provenance.json"))
    assert "raised_during_write" not in [w["code"] for w in loaded["warnings"]]


# ---------------------------------------------------------------------------
# a key to the pictures
# ---------------------------------------------------------------------------

def test_the_tables_say_which_colour_each_row_is_drawn_in(fixture):
    """The pictures had no key: a tester matched colours to phases by
    counting pixels in the PNG against ``frac_of_scan``. That works until two
    phases have similar areas, which is the ordinary case."""
    from backend.api.services.phase_map_store import (
        palette_hex_for_state, region_color_hex,
    )
    state, _ = fixture
    tables = _build(fixture)

    assert "phase_map_colour_hex" in tables.phases.columns
    assert "phase_map_colour_hex" in tables.regions.columns
    assert "region_map_colour_hex" in tables.regions.columns

    # The SAME functions the renderers call, so the key cannot disagree with
    # the image it describes.
    palette = list(palette_hex_for_state(state, None))
    for row in tables.phases.rows:
        idx = row["phase_index"]
        expected = X._NO_DATA_HEX if idx < 0 else palette[idx]
        assert row["phase_map_colour_hex"] == expected, row["phase_name"]
    for row in tables.regions.rows:
        assert row["region_map_colour_hex"] == \
            region_color_hex(row["region_id"])
        assert row["phase_map_colour_hex"] == \
            palette[row["phase_final_index"]]


def test_two_regions_of_one_phase_share_a_phase_colour_and_not_a_region_one(
        fixture):
    """Which is exactly why regions.csv needs BOTH columns.

    A region is one colour in region_map.png (its own) and a different colour
    in phase_map.png (its phase's, shared with every other region carrying
    that phase). One column could only answer one of the two questions, and a
    reader holding the other picture would match the wrong name to the
    colour.
    """
    tables = _build(fixture)
    rows = {r["region_id"]: r for r in tables.regions.rows}
    assert rows[1]["phase_final"] == rows[2]["phase_final"] == "Si.cif"
    assert rows[1]["phase_map_colour_hex"] == rows[2]["phase_map_colour_hex"]
    assert rows[1]["region_map_colour_hex"] != rows[2]["region_map_colour_hex"]


def test_the_no_data_grey_is_recorded_where_no_row_can_carry_it(fixture):
    """Both renderers paint unassigned pixels (60, 60, 60). It is stated in
    the provenance because a legend that names the wrong grey tells the
    reader the hole in their map is a phase."""
    tables = _build(fixture)
    c = tables.provenance["colours"]
    assert c["available"] is True
    assert c["no_data_hex"] == "#3c3c3c"
    assert "unassigned" in c["no_data_note"]
    assert "phase_map_store" in c["source"]
    row = _row(tables.phases, phase_name="unclassified")
    assert row["phase_map_colour_hex"] == "#3c3c3c"


def test_a_colour_override_reaches_the_column_and_is_recorded(fixture):
    """A user who recoloured a phase gets that phase in that colour — in the
    picture AND in the key to it."""
    tables = _build(fixture, colour_overrides={"Si.cif": "#ff00ff"})
    assert _row(tables.phases,
                phase_name="Si.cif")["phase_map_colour_hex"] == "#ff00ff"
    assert tables.provenance["colours"]["overrides_applied"] == ["Si.cif"]


def test_the_colours_are_there_even_when_the_picture_is_turned_off(fixture):
    """The colours are a fact about the map the app is showing, not about
    which artefacts were requested. An empty column would say the phase has
    no colour rather than that the reader turned the picture off."""
    tables = _build(fixture, artefacts={
        "phases": True, "regions": True, "particles": True,
        "definitions": True, "xlsx": False, "map_png": False})
    assert tables.images == {}
    assert _row(tables.phases, phase_name="Si.cif")["phase_map_colour_hex"]


def test_a_palette_that_cannot_be_read_warns_and_leaves_the_cells_empty(
        fixture, monkeypatch):
    """Never a plausible-looking hex. Printing somebody else's colour beside
    this row's name is the one failure a legend must not have."""
    import backend.api.services.phase_map_store as store

    def _boom(*_a, **_k):
        raise RuntimeError("no palette")

    monkeypatch.setattr(store, "palette_hex_for_state", _boom)
    tables = _build(fixture)
    assert all(r["phase_map_colour_hex"] == "" for r in tables.phases.rows)
    assert tables.provenance["colours"]["available"] is False
    assert [w for w in tables.warnings if w["code"] == "legend_unavailable"]
    # ...and not one number moved.
    assert _row(tables.phases, phase_name="Si.cif")["n_px"] == \
        N_REGION1 + N_REGION2


# ---------------------------------------------------------------------------
# the two small ones
# ---------------------------------------------------------------------------

def test_the_min_particle_px_default_travels_with_the_value(fixture):
    """So "was this run configured, or did it inherit?" is answerable.

    A tester reported the scripted and the clicked export disagreeing here.
    MEASURED on this repo they do not — ``ExportOptions``, ``ExportRequest``
    and the dialog's ``useState(0)`` are all 0 — but agreement today is
    exactly the kind of thing that drifts silently, so the default travels
    with the value and a future divergence shows up in the file.
    """
    tables = _build(fixture)
    parts = tables.provenance["particles"]
    assert parts["min_particle_px"] == 0
    assert parts["min_particle_px_default"] == 0
    assert parts["min_particle_px_used_default"] is True

    chosen = _build(fixture, min_particle_px=5)
    parts = chosen.provenance["particles"]
    assert parts["min_particle_px"] == 5
    assert parts["min_particle_px_default"] == 0
    assert parts["min_particle_px_used_default"] is False


def test_the_recorded_default_is_the_dataclass_default_and_not_a_copy(fixture):
    """Two places holding one fact is this module's own worst defect class.

    Read off ``ExportOptions`` itself, so a change to the default cannot
    leave the provenance asserting the old one.
    """
    tables = _build(fixture)
    assert tables.provenance["particles"]["min_particle_px_default"] == \
        X.ExportOptions().min_particle_px


def test_the_coordinate_order_claim_names_where_it_applies(fixture):
    """It used to assert "every flat index in this export is
    row * n_cols + col" — and no table carries a flat index, so it documented
    an absent thing."""
    tables = _build(fixture)
    order = tables.provenance["coordinates"]["order"]
    assert "labels.npz" in order
    assert "NO TABLE IN THIS EXPORT CARRIES A FLAT INDEX" in order

    # ...and the claim it now makes is true of the artefact it names.
    assert tables.labels["phase_id"].shape == (N_ROWS, N_COLS)
    flat = tables.labels["phase_id"].ravel()
    row, col = 3, 7
    assert flat[row * N_COLS + col] == tables.labels["phase_id"][row, col]

    # ...and the columns it points the reader at instead really are there,
    # while the flat index it no longer claims really is absent.
    for c in ("centroid_row", "centroid_col", "bbox_row_min"):
        assert c in tables.particles.columns
    assert "row" in tables.pixel_columns and "col" in tables.pixel_columns
    for table in (tables.phases, tables.regions, tables.particles):
        assert not [c for c in table.columns if "flat" in c.lower()]


# ---------------------------------------------------------------------------
# the picture
# ---------------------------------------------------------------------------

def test_the_map_pngs_are_written_by_default(tmp_path, fixture):
    """Opt-in-by-default. The reader who asked for this had already missed a
    wrong headline number for want of a picture; an artefact she would have
    had to go and switch on is the same wall with an extra checkbox."""
    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    names = {f.name: f for f in res.files}
    assert "phase_map.png" in names
    assert "region_map.png" in names
    assert names["phase_map.png"].bytes > 0
    # `rows` on an image is its pixel count, the number that says at a glance
    # whether it covers the whole raster.
    assert names["phase_map.png"].rows == N_PX

    from PIL import Image
    with Image.open(Path(res.folder) / "phase_map.png") as im:
        im.load()
        assert im.format == "PNG"
        assert im.size == (N_COLS, N_ROWS)


def test_the_exported_png_is_the_picture_the_app_draws(tmp_path, fixture):
    """Byte-identical to ``phase_map_store``'s own renderer, overrides and all.

    Not a second drawing. A picture in the folder that coloured a phase
    differently from the screen would be worse than no picture: a reader
    comparing the two would conclude one of them is a different map.
    """
    import base64
    from backend.api.services import phase_map_store as pms

    state, at_maps = fixture
    overrides = {"Si": "#ff00ff"}
    tables = X.build_tables(
        state, at_maps, geometry=_geom(),
        options=X.ExportOptions(hash_source=False, colour_overrides=overrides))
    assert tables.images["phase_map.png"] == base64.b64decode(
        pms.render_phase_map_to_base64(state, overrides))
    assert tables.images["region_map.png"] == base64.b64decode(
        pms.render_region_map_to_base64(state))
    # and the overrides are not decorative: a different colour, a different file
    assert tables.images["phase_map.png"] != base64.b64decode(
        pms.render_phase_map_to_base64(state, {"Si": "#00ff00"}))


def test_a_pixel_mode_map_gets_a_phase_picture_and_no_region_one(fixture):
    """``render_region_map_to_base64`` raises without regions, and a warning
    saying "the region picture is missing" about a map that has none would be
    noise rather than information."""
    state, at_maps = fixture
    state.region_grid = None
    state.region_phase = []
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert "phase_map.png" in tables.images
    assert "region_map.png" not in tables.images
    assert not [w for w in tables.warnings if w["code"] == "map_png_failed"]


def test_turning_the_picture_off_leaves_the_numbers_alone(tmp_path, fixture):
    opts = X.ExportOptions(hash_source=False, artefacts={
        "phases": True, "regions": True, "particles": True,
        "definitions": True, "xlsx": False, "map_png": False})
    tables = X.build_tables(fixture[0], fixture[1], geometry=_geom(),
                            options=opts)
    assert tables.images == {}
    res = X.write_export(tables, tmp_path, options=opts)
    assert {f.name for f in res.files} == {
        "phases.csv", "regions.csv", "particles.csv", "definitions.csv",
        # caption.txt is written whenever there is a summary to write —
        # not an artefact toggle. An export without a caption is the one
        # whose numbers get quoted without their conditions.
        "provenance.json", "caption.txt"}


def test_a_failing_render_costs_the_picture_and_never_the_numbers(
        tmp_path, fixture, monkeypatch):
    from backend.api.services import phase_map_store as pms

    def _boom(*_a, **_kw):
        raise RuntimeError("no Pillow here")

    monkeypatch.setattr(pms, "render_phase_map_to_base64", _boom)
    tables = _build(fixture)
    assert "phase_map.png" not in tables.images
    assert any(w["code"] == "map_png_failed" for w in tables.warnings)

    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    names = {f.name for f in res.files}
    assert "phases.csv" in names and "provenance.json" in names
    # the region picture still made it - a partial failure is partial
    assert "region_map.png" in names
    assert any(w["code"] == "map_png_failed" for w in res.warnings)


def test_the_png_pixel_count_is_read_from_the_header(fixture):
    tables = _build(fixture)
    assert X._png_pixel_count(tables.images["phase_map.png"]) == N_PX
    assert X._png_pixel_count(b"not a png at all") == 0
    assert X._png_pixel_count(b"") == 0
# ---------------------------------------------------------------------------
# a particle's provenance is that particle's
# ---------------------------------------------------------------------------

def test_a_particle_in_a_painted_region_does_not_claim_to_have_been_painted(fixture):
    """THE defect, stated as a test, and the one to break first.

    A QA tester measured a real export: 110 of 167 particles read
    ``assignment_source = mixed`` while carrying ``n_px_hand = 0``, with
    exactly ONE particle in the whole file actually holding a painted pixel
    — and the file's own legend defines ``mixed`` as "SOME pixels were
    painted". 110 rows of an auditor's table asserted a human decision that
    did not happen.

    Region 1 is in two disconnected pieces. Paint the 4-pixel piece and
    nothing else: that piece must read ``mixed`` and the untouched 9-pixel
    piece must not, however the region as a whole is graded.
    """
    state, at_maps = fixture
    flat_region = state.region_grid.ravel()
    rows = np.arange(N_PX) // N_COLS
    small_piece = (flat_region == 1) & (rows >= 6)          # the 4 px block
    assert small_piece.sum() == 4

    locked = np.zeros(N_PX, dtype=bool)
    locked[small_piece] = True
    state.locked_mask = locked.reshape(N_ROWS, N_COLS)

    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))

    # The region genuinely is mixed: 4 of its 13 pixels were painted.
    assert _row(tables.regions, region_id=1)["assignment_source"] == "mixed"

    pieces = {r["n_px"]: r for r in tables.particles.rows if r["region_id"] == 1}
    assert sorted(pieces) == [4, 9]
    assert pieces[4]["n_px_hand"] == 4
    assert pieces[4]["assignment_source"] == "hand"     # all of ITS pixels
    assert pieces[9]["n_px_hand"] == 0
    # The pre-fix value was "mixed" — the region's grade, on a particle no
    # hand ever touched. Spelled out so a revert fails instead of merely
    # looking plausible.
    assert pieces[9]["assignment_source"] != "mixed"
    assert pieces[9]["assignment_source"] == "auto"


def test_no_particle_anywhere_says_mixed_without_a_painted_pixel(fixture):
    """The invariant behind the count, over every particle of the map."""
    state, at_maps = fixture
    tables = _build(fixture)
    for r in tables.particles.rows:
        if r["assignment_source"] in ("hand", "mixed"):
            assert r["n_px_hand"] > 0, r
        if r["assignment_source"] == "hand":
            assert r["n_px_hand"] == r["n_px"], r


def test_a_particle_still_inherits_the_decision_that_was_its_regions(fixture):
    """What DOES propagate, and why it is not the same thing.

    ``definition`` / ``auto`` / ``untracked`` describe how the phase NAME
    was arrived at, and under region basis that decision genuinely was the
    region's — the particle was never named individually. Point region 1 at
    the wrong phase, as a rename would, and its particles have to inherit
    ``untracked`` even though nothing was painted anywhere.
    """
    state, at_maps = fixture
    state.region_phase[1] = 0                  # user says Al, chemistry says Si
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert _row(tables.regions, region_id=1)["assignment_source"] == "untracked"
    for r in tables.particles.rows:
        if r["region_id"] == 1:
            assert r["n_px_hand"] == 0
            assert r["assignment_source"] == "untracked"


def test_a_regions_painting_cannot_leak_back_through_the_inheritance(fixture):
    """The subtle half: the inherited part is recomputed with n_hand = 0.

    Region 0 owns 83 pixels and 4 of them are painted, none of them inside
    the particle being checked. Inheriting the region's grade verbatim gives
    ``mixed``; inheriting its DECISION gives ``auto``. The difference is the
    whole defect.
    """
    tables = _build(fixture)                    # the fixture paints 4 px of region 0
    assert _row(tables.regions, region_id=0)["assignment_source"] == "mixed"
    for r in tables.particles.rows:
        if r["region_id"] == 0 and r["n_px_hand"] == 0:
            assert r["assignment_source"] == "auto"


def test_dissent_gets_its_own_column_rather_than_riding_in_the_provenance_one(fixture):
    """``auto`` beside ``phase_agrees = False`` reproduced the wrong inference.

    A second tester came at the same problem from the other side: on a clean
    scan the pair fires on a large fraction of the rows, and a reader with
    only the CSV concludes a human moved them. The fact is worth exporting —
    it is exactly the row worth finding — so it gets a column whose NAME
    says what it is.
    """
    state, at_maps = fixture
    state.region_phase[1] = 0                  # chemistry says Si, map says Al
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert "chemistry_dissents" in tables.particles.columns
    for r in tables.particles.rows:
        # One fact, two spellings, and they can never disagree.
        assert r["chemistry_dissents"] == (r["phase_agrees"] is False)
    dissenting = [r for r in tables.particles.rows if r["chemistry_dissents"]]
    assert dissenting, "region 1 was pointed at the wrong phase"
    # And dissent alone never becomes a provenance claim about the particle.
    for r in dissenting:
        assert r["n_px_hand"] == 0
        assert r["assignment_source"] != "mixed"


def test_the_note_describes_what_the_code_now_does(fixture):
    """The worked example in the provenance was wrong about its own result.

    It promised a dissenting particle "will read assignment_source = auto
    with phase_agrees = False". Measured over the 97 dissenting particles of
    a real export: auto 7, mixed 81, untracked 9 — the documented example
    held for 7 %. A note and a behaviour that disagree are worse than
    neither, because the note is what a reader trusts.
    """
    prov = _build(fixture).provenance["particles"]
    note = prov["assignment_source_note"]
    assert "n_px_hand" in note
    assert "110 of 167" in note                  # the measurement it fixes
    assert "chemistry_dissents_note" in prov
    assert "not a provenance claim" in prov["chemistry_dissents_note"].lower()
    # The old sentence claimed a specific pairing that was false 93 % of the
    # time. It must not survive anywhere in the record.
    assert "will read assignment_source = auto" not in note


# ---------------------------------------------------------------------------
# which phases never competed, and which library said so
# ---------------------------------------------------------------------------

def test_a_phase_dropped_for_an_unmeasured_element_is_named_with_the_element():
    """``candidates_for`` is a silent subset filter, and silence was the bug.

    A QA tester measured her library at 30 entries, 22 competing and 8
    dropped — seven for magnesium and one for nickel — with nothing in the
    export to say so. Here: a two-element scan against a library containing
    a copper phase.
    """
    lib = {e.key: e for e in ENTRIES}
    identity, absent = X.library_report(
        ENTRIES[:2], ["Al", "Si"], library=lib, library_path="lib.xlsx")
    assert identity["n_entries"] == 3
    assert identity["n_participating"] == 2
    dropped = {d["key"]: d for d in absent}
    assert set(dropped) == {"Cu2Si.cif"}
    assert dropped["Cu2Si.cif"]["reason"] == "element_not_measured"
    assert dropped["Cu2Si.cif"]["missing_elements"] == ["Cu"]


def test_a_phase_the_run_asked_for_and_the_library_lacks_is_a_different_reason():
    """The two cases the tester could not tell apart.

    Her beta was in the library under a different filename and DID compete;
    her gamma-Al3FeSi was not in the library at all. One is fixed by mapping
    an element or widening a selection, the other by adding a CIF, and an
    export that reports both as "absent" sends her to the wrong place.
    """
    lib = {e.key: e for e in ENTRIES}
    _identity, absent = X.library_report(
        ENTRIES[:1], ["Al", "Si"], library=lib, library_path="lib.xlsx",
        requested_keys=["Al.cif", "Si.cif", "gamma-Al3FeSi.cif"])
    by_key = {d["key"]: d["reason"] for d in absent}
    assert by_key["gamma-Al3FeSi.cif"] == "not_in_library"
    # Si was in the library and every element of it was measured; it simply
    # was not among the entries that competed while phase_keys named it.
    assert by_key["Si.cif"] == "excluded_upstream"
    assert by_key["Cu2Si.cif"] == "element_not_measured"


def test_a_library_entry_the_run_narrowed_out_says_so_rather_than_blaming_chemistry():
    lib = {e.key: e for e in ENTRIES}
    _identity, absent = X.library_report(
        ENTRIES[:1], ["Al", "Si", "Cu"], library=lib, library_path="lib.xlsx",
        requested_keys=["Al.cif"])
    by_key = {d["key"]: d["reason"] for d in absent}
    assert by_key["Si.cif"] == "not_requested"
    # Every element measured, so the chemistry reason would be a lie.
    assert by_key["Cu2Si.cif"] == "not_requested"


def test_the_library_identity_moves_when_the_file_does_and_says_which_hash(tmp_path):
    """Three testers asked for the hash independently, and one said why.

    ``phases_participating`` already exposes a stoichiometry edit by diffing
    two exports. It does NOT expose a lattice-parameter-only edit — that
    changes nothing this module parses — and that is the edit he cares about
    between heat treatments. Only the FILE hash can see it, so both hashes
    are exported and they answer different questions.
    """
    lib_a = tmp_path / "crystal_database.xlsx"
    lib_a.write_bytes(b"a database, version one")
    lib = {e.key: e for e in ENTRIES}
    ident_a, _ = X.library_report(ENTRIES, ["Al", "Si", "Cu"], library=lib,
                                  library_path=str(lib_a))
    assert ident_a["sha256"] == \
        __import__("hashlib").sha256(lib_a.read_bytes()).hexdigest()
    assert ident_a["bytes"] == lib_a.stat().st_size
    assert ident_a["mtime"]

    # An edit this export never parses: the bytes move, the entries do not.
    lib_a.write_bytes(b"a database, version two -- one lattice parameter")
    ident_b, _ = X.library_report(ENTRIES, ["Al", "Si", "Cu"], library=lib,
                                  library_path=str(lib_a))
    assert ident_b["sha256"] != ident_a["sha256"]
    assert ident_b["entries_digest"] == ident_a["entries_digest"]

    # A chemistry edit moves both.
    changed = list(ENTRIES[:2]) + [_entry("Cu2Si.cif", {"Cu": 50.0, "Si": 50.0})]
    ident_c, _ = X.library_report(
        changed, ["Al", "Si", "Cu"], library={e.key: e for e in changed},
        library_path=str(lib_a))
    assert ident_c["entries_digest"] != ident_a["entries_digest"]


def test_an_unreadable_library_says_so_instead_of_reporting_none_dropped():
    """"no phases were dropped" and "we could not check" are different."""
    identity, absent = X.library_report(ENTRIES, ["Al", "Si"], library={},
                                        library_path="nowhere.xlsx")
    assert absent == []
    assert identity["n_entries"] == 0
    assert "could not be read" in identity["phases_absent_note"]
    # Not silently reported as a clean bill of health.
    assert "n_absent" not in identity


def test_the_export_carries_the_library_block_and_warns_about_the_drop(fixture):
    """End to end, through ``build_tables``, with an injected library."""
    state, at_maps = fixture
    lib = {e.key: e for e in ENTRIES}
    opts = X.ExportOptions(hash_source=False, library=lib,
                           library_path="lib.xlsx")
    tables = X.build_tables(state, at_maps, geometry=_geom(), options=opts)
    prov = tables.provenance
    assert prov["library"]["source"] == "caller"
    assert prov["library"]["n_entries"] == 3
    assert [d["key"] for d in prov["phases_absent"]] == []   # Cu2Si competes
    # Now hide copper from the measurement so the phase is genuinely dropped.
    state.phase_entries = ENTRIES[:2]
    tables = X.build_tables(state, at_maps, geometry=_geom(), options=opts)
    assert [d["reason"] for d in tables.provenance["phases_absent"]] == \
        ["element_not_measured"]
    assert any(w["code"] == "phases_dropped_unmeasured_element"
               for w in tables.warnings)


# ---------------------------------------------------------------------------
# how the spectra were acquired
# ---------------------------------------------------------------------------

def test_the_acquisition_header_reaches_the_export(fixture):
    """Present in the source, absent from the export, and not even declared."""
    state, at_maps = fixture
    geom = X.ExportGeometry(
        n_rows=N_ROWS, n_cols=N_COLS, step_x_um=STEP, step_y_um=STEP,
        source_path="fixture.h5oina",
        acquisition={"Beam Voltage": 20.0, "Working Distance": 16.0366,
                     "Process Time": 4, "Detector Serial Number": "UVA11191",
                     "Window Type": "SATW", "Something Aztec Adds Later": 7})
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    acq = tables.provenance["acquisition"]
    assert acq["available"] is True
    assert acq["source"] == "caller"
    assert acq["beam_voltage_kv"] == pytest.approx(20.0)
    assert acq["working_distance_mm"] == pytest.approx(16.0366)
    assert acq["process_time"] == 4
    assert acq["detector_serial"] == "UVA11191"
    assert acq["window_type"] == "SATW"
    # A key this exporter has never heard of survives rather than vanishing.
    assert acq["raw"]["Something Aztec Adds Later"] == 7
    assert not any(w["code"] == "no_acquisition_header" for w in tables.warnings)


def test_a_missing_header_is_declared_rather_than_defaulted(fixture):
    state, at_maps = fixture
    tables = X.build_tables(
        state, at_maps, geometry=_geom(),
        options=X.ExportOptions(hash_source=False,
                                read_acquisition_header=False))
    acq = tables.provenance["acquisition"]
    assert acq["available"] is False
    assert acq["source"] == "not requested"
    assert "beam_voltage_kv" not in acq
    assert any(w["code"] == "no_acquisition_header" for w in tables.warnings)
    # And the omission is visible in the record, not only in the warning.
    assert tables.provenance["interaction_volume"]["range_um"] is None
    assert "acquisition.available" in \
        tables.provenance["interaction_volume"]["note"]


def test_the_kanaya_okayama_range_is_the_number_the_tester_computed():
    """4.19 um in aluminium at 20 kV, against a 0.5 um step.

    Hand-checked: R = 0.0276 * 26.98 * 20^1.67 / (13^0.89 * 2.70).
    """
    assert X.kanaya_okayama_range_um("Al", 20.0) == pytest.approx(4.19, abs=0.01)
    # Monotonic in voltage and heavier elements stop the beam sooner.
    assert X.kanaya_okayama_range_um("Al", 30.0) > \
        X.kanaya_okayama_range_um("Al", 20.0)
    assert X.kanaya_okayama_range_um("Fe", 20.0) < \
        X.kanaya_okayama_range_um("Al", 20.0)
    # Never guessed: an unknown element or an unusable voltage is empty.
    assert X.kanaya_okayama_range_um("Unobtainium", 20.0) is None
    assert X.kanaya_okayama_range_um("Al", 0.0) is None
    assert X.kanaya_okayama_range_um("Al", None) is None


def test_the_interaction_volume_is_compared_against_the_particles(fixture):
    """Prose becomes a number: how many particles are below the range."""
    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=STEP,
                            step_y_um=STEP, source_path="fixture.h5oina",
                            acquisition={"Beam Voltage": 20.0})
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    iv = tables.provenance["interaction_volume"]
    assert iv["matrix_element"] == "Al"
    assert iv["range_um"] == pytest.approx(4.19, abs=0.01)
    assert iv["step_um"] == pytest.approx(STEP)
    assert iv["range_over_step"] == pytest.approx(4.19 / STEP, abs=0.05)

    # Counted independently off the particle rows, on the same column.
    ecds = [r["ecd_um"] for r in tables.particles.rows]
    expected = sum(1 for v in ecds if v < iv["range_um"])
    assert iv["n_particles_measured"] == len(ecds)
    assert iv["n_particles_smaller_than_range"] == expected
    assert iv["frac_particles_smaller_than_range"] == \
        pytest.approx(expected / len(ecds))


def test_the_real_scans_header_reads_back_including_its_subgroup():
    """``Stage Position`` is a GROUP, so a flat ``group[k][()]`` raises on it.

    Skipped when the file is not on this machine; when it is, this is the
    only test that proves the reader against a real vendor file rather than
    against a dict somebody wrote to match it.
    """
    path = (ROOT / "Test_data" /
            "EBSD_SampleB_extrusion_withPattern Sample_B Arbeitsbereich 1 "
            "Elementverteilungsdaten 1.h5oina")
    if not path.is_file():
        pytest.skip("the SampleB scan is not on this machine")
    header = X.read_acquisition_header(str(path))
    assert header is not None
    assert header["Beam Voltage"] == pytest.approx(20.0)
    assert header["Detector Serial Number"] == "UVA11191"
    assert header["Acquisition Date"] == "2026-01-19T09:04:03"
    assert isinstance(header["Stage Position"], dict)
    assert "Rotation" in header["Stage Position"]


def test_the_header_is_read_from_the_source_when_the_caller_supplies_none(
        tmp_path, fixture):
    """The path every current caller takes, and it must be exercised.

    ``routes/eds_export.py`` builds its ``ExportGeometry`` without an
    acquisition header, so the fallback read from ``source_path`` is what
    actually runs in the app. A mutation that made that read return None was
    caught by no test until this one existed -- every other acquisition test
    hands the dict in directly.
    """
    h5py = pytest.importorskip("h5py")
    src = tmp_path / "tiny.h5oina"
    with h5py.File(src, "w") as fh:
        hdr = fh.create_group("1/EDS/Header")
        hdr.create_dataset("Beam Voltage", data=np.array([15.0], dtype="f4"))
        hdr.create_dataset("Detector Serial Number",
                           data=np.array([b"SN-1"], dtype=object))
        # A GROUP, not a dataset -- the shape that makes a flat read raise.
        stage = hdr.create_group("Stage Position")
        stage.create_dataset("Tilt", data=np.array([70.0], dtype="f4"))

    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=STEP,
                            step_y_um=STEP, source_path=str(src))
    assert geom.acquisition is None
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(hash_source=False))
    acq = tables.provenance["acquisition"]
    assert acq["available"] is True
    assert acq["source"] == "source file"
    assert acq["beam_voltage_kv"] == pytest.approx(15.0)
    assert acq["detector_serial"] == "SN-1"
    assert acq["raw"]["Stage Position"]["Tilt"] == pytest.approx(70.0)
    # And the interaction volume follows from it without any other input.
    iv = tables.provenance["interaction_volume"]
    assert iv["beam_voltage_kv"] == pytest.approx(15.0)
    assert iv["range_um"] == pytest.approx(
        X.kanaya_okayama_range_um("Al", 15.0))


def test_no_source_no_header_and_no_crash():
    assert X.read_acquisition_header(None) is None
    assert X.read_acquisition_header("") is None
    assert X.read_acquisition_header("does-not-exist.h5oina") is None


# ---------------------------------------------------------------------------
# the edit log, and the two fields that had no note
# ---------------------------------------------------------------------------

def test_the_edit_log_reaches_the_file_it_claims_to_be_complete_about(fixture):
    """``log_retained: 5`` beside "every edit is listed", listing nothing.

    That sentence was false about the file it appeared in. The store keeps
    the log and ``edit_summary()`` returns it already JSON-shaped; it was
    dropped on the way out.
    """
    state, at_maps = fixture

    class _Stated:
        """A state that answers `edit_summary` the way the store does."""
        def __init__(self, inner):
            self._inner = inner

        def __getattr__(self, name):
            return getattr(self._inner, name)

        def edit_summary(self):
            return {
                "tracked": True,
                "counts": {"merges": 1, "paints": 0, "regions_named": 1,
                           "splits": 0, "grows": 0, "edge_snaps": 0,
                           "phase_replacements": 0},
                "n_edits": 2, "log_retained": 2, "log_truncated": False,
                "log_cap": 500, "region_grid_edited": True,
                "log": [
                    {"op": "merge", "detail": {"keep_id": 2, "drop_id": 3,
                                               "pixels_moved": 66}},
                    {"op": "name_region", "detail": {"region_id": 1}},
                ],
                "note": "Complete: every edit since classification is listed.",
            }

    tables = X.build_tables(_Stated(state), at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    h = tables.provenance["hand_edits"]
    assert h["log_retained"] == 2
    assert len(h["log"]) == 2                       # it was [] — i.e. absent
    assert h["log"][0]["op"] == "merge"
    assert h["log"][0]["detail"]["pixels_moved"] == 66
    assert h["log_cap"] == 500
    assert "oldest" in h["log_note"]
    # The log is what tells "a rename happened" apart from "the replay
    # produced no answer", which `untracked` alone cannot.
    assert any(e["op"] == "name_region" for e in h["log"])


def test_an_untracked_map_has_a_null_log_and_not_an_empty_one(fixture):
    h = _build(fixture).provenance["hand_edits"]
    assert h["log"] is None
    assert "assert that nothing happened" in h["log_note"]


def test_a_renumbering_edit_raises_a_warning_and_not_only_a_boolean(fixture):
    """A merge renumbers particle ids: 58 of 135 changed at 1.000 overlap.

    ``region_grid_edited`` was a bare boolean with its meaning in a Python
    comment. The thing most likely to break a figure has to be a warning,
    because warnings are what the UI shows.
    """
    state, at_maps = fixture

    class _Edited:
        def __init__(self, inner):
            self._inner = inner

        def __getattr__(self, name):
            return getattr(self._inner, name)

        def edit_summary(self):
            return {"tracked": True, "counts": {"merges": 1},
                    "n_edits": 1, "log_retained": 1, "log_truncated": False,
                    "log_cap": 500, "region_grid_edited": True,
                    "log": [{"op": "merge", "detail": {}}], "note": "x"}

    tables = X.build_tables(_Edited(state), at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    codes = [w["code"] for w in tables.warnings]
    assert "region_ids_renumbered" in codes
    msg = next(w["message"] for w in tables.warnings
               if w["code"] == "region_ids_renumbered")
    assert "RENUMBER" in msg
    h = tables.provenance["hand_edits"]
    assert h["region_grid_edited"] is True
    assert "RENUMBER" in h["region_grid_edited_note"]

    # And an unedited map stays quiet.
    assert not [w for w in _build(fixture).warnings
                if w["code"] == "region_ids_renumbered"]


def test_paints_zero_beside_painted_pixels_is_reconciled_in_one_sentence(fixture):
    """Both are true under their own definitions; nothing said so.

    ``paints`` counts operations SINCE THE LAST CLASSIFICATION and a
    re-classify resets it; ``pixels_painted`` counts the locked pixels,
    which a re-classify deliberately carries across.
    """
    note = _build(fixture).provenance["hand_edits"]["paints_note"]
    assert "re-classify" in note
    assert "pixels_painted" in note and "paints" in note


# ---------------------------------------------------------------------------
# the caption, and names a report can carry
# ---------------------------------------------------------------------------

def test_the_summary_is_in_the_record_and_in_the_folder(tmp_path, fixture):
    tables = _build(fixture)
    line = tables.provenance["summary"]
    assert line.endswith(".")
    assert "10x10 px at 0.5 um step" in line
    assert "3 regions -> 2 phases" in line
    assert "semi-quantitative" in line

    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    caption = (Path(res.folder) / "caption.txt").read_text(encoding="utf-8")
    assert caption.strip() == line
    assert "caption.txt" in {f.name for f in res.files}


def test_an_unavailable_clause_is_omitted_and_never_printed_as_zero(fixture):
    """The frontend's version of this line shipped a Number(null) === 0 bug.

    A step of 0.0 um, a count of 0 and a missing value are three different
    things. Nothing here may default; a clause with no value is dropped.
    """
    state, at_maps = fixture
    geom = X.ExportGeometry(n_rows=N_ROWS, n_cols=N_COLS, step_x_um=None,
                            step_y_um=None, source_path=None)
    tables = X.build_tables(state, at_maps, geometry=geom,
                            options=X.ExportOptions(
                                hash_source=False,
                                read_acquisition_header=False))
    line = tables.provenance["summary"]
    assert "um step" not in line                  # no step: no step clause
    assert "kV" not in line                       # no header: no voltage
    assert "interaction range" not in line
    assert "0 um" not in line and "0 kV" not in line
    # The clauses that DO have values are still there.
    assert "10x10 px" in line
    assert "semi-quantitative" in line
    # A real count is printed, not dropped: the fixture paints 4 pixels.
    assert f"{N_HAND} hand-painted pixels" in line
    # And a real zero is a zero rather than an omission.
    state.locked_mask = None
    bare = X.build_tables(state, at_maps, geometry=geom,
                          options=X.ExportOptions(
                              hash_source=False,
                              read_acquisition_header=False))
    assert "no hand-painted pixels" in bare.provenance["summary"]


def test_a_phase_gets_a_name_a_figure_legend_can_carry(fixture):
    """``sd_0302719.cif`` and ``Mn4.512Al127.296...`` are neither of them names."""
    ugly = _entry("sd_0302719.cif", {"Al": 74.0, "Fe": 11.3, "Si": 9.7, "Mn": 2.6},
                  formula="Mn\u2084.\u2085\u2081\u2082Al\u2081\u2082\u2087")
    assert X.phase_label(ugly) == "Al-Fe-Si-Mn (sd_0302719)"
    assert X.formula_ascii(ugly.formula) == "Mn4.512Al127"

    # A filename that already names the phase is left alone, minus the
    # database bookkeeping nobody puts in a legend.
    assert X.phase_label(_entry("Al6Fe_mp-570001_symmetrized.cif",
                                {"Al": 85.7, "Fe": 14.3})) == "Al6Fe"
    assert X.phase_label(_entry("beta-AlFeSi.cif", {"Al": 84.6, "Fe": 15.4})) \
        == "beta-AlFeSi"
    assert X.phase_label(None) == ""


def test_the_readable_names_are_columns_and_not_only_provenance(fixture):
    tables = _build(fixture)
    assert "phase_label" in tables.phases.columns
    assert "formula_ascii" in tables.phases.columns
    assert "phase_final_label" in tables.regions.columns
    assert "phase_label" in tables.particles.columns
    r = _row(tables.phases, phase_name="Al.cif")
    assert r["phase_label"] == "Al"
    assert _row(tables.regions, region_id=1)["phase_final_label"] == "Si"
    assert {p["phase_label"] for p in tables.provenance["phases_participating"]} \
        == {"Al", "Si", "Cu2Si"}


# ---------------------------------------------------------------------------
# the smaller asks
# ---------------------------------------------------------------------------

def test_third_place_is_exported_because_a_margin_needs_it(fixture):
    """A 0.15 margin means one thing over a distant third and another over
    a near one, and the second case is a three-way tie reported as a
    decision."""
    tables = _build(fixture)
    for table in (tables.regions, tables.particles):
        assert "third_phase" in table.columns
        assert "third_score" in table.columns
    r = _row(tables.regions, region_id=1)
    assert r["third_score"] is not None
    assert r["runner_up_score"] >= r["third_score"]
    assert r["third_phase"] not in (r["phase_auto"], r["runner_up_phase"])

    # A two-candidate library has no third place, and must not repeat one.
    state, at_maps = fixture
    state.phase_entries = ENTRIES[:2]
    state.region_phase = [0, 1, 1]
    two = X.build_tables(state, at_maps, geometry=_geom(),
                         options=X.ExportOptions(hash_source=False))
    r2 = _row(two.regions, region_id=1)
    assert r2["runner_up_phase"] != ""
    assert r2["third_phase"] == ""
    assert r2["third_score"] is None


def test_a_phase_row_summarises_the_particle_sizes_behind_its_count(fixture):
    """``n_particles`` without a size distribution is half a sentence."""
    tables = _build(fixture)
    r = _row(tables.phases, phase_name="Si.cif")
    sizes = [p["area_um2"] for p in tables.particles.rows
             if p["phase_name"] == "Si.cif"]
    assert sizes
    assert r["particle_area_mean_um2"] == pytest.approx(float(np.mean(sizes)))
    assert r["particle_area_max_um2"] == pytest.approx(max(sizes))
    assert r["particle_area_sd_um2"] == pytest.approx(float(np.std(sizes)))


def test_one_particle_has_no_spread_and_says_so_rather_than_zero(fixture):
    """"0" would read as "they are all the same size"."""
    tables = _build(fixture)
    for r in tables.phases.rows:
        if r.get("n_particles") == 1:
            assert r["particle_area_sd_um2"] is None
            assert r["particle_area_mean_um2"] is not None
            break
    else:
        pytest.skip("no single-particle phase in this fixture")


def test_counting_particles_of_the_matrix_is_flagged_as_a_category_error(fixture):
    """A tester read "143 Al particles" and called it thesis material.

    They are connected components of a CONTINUOUS phase. The rows are not
    wrong; calling them particles is, and only the file can say so.

    The shared fixture cannot show it: its matrix is one connected piece, so
    "1 particle" is not yet a misleading COUNT. Cut it in two with a band of
    the silicon region and the matrix starts reporting a count, which is the
    number that ends up in a thesis.
    """
    # First, the negative: a matrix in one piece is not a misleading count
    # and must stay quiet. Checked BEFORE the mutation below, because the
    # fixture state is mutated in place.
    assert not [w for w in _build(fixture).warnings
                if w["code"] == "matrix_counted_as_particles"]

    state, at_maps = fixture
    region = state.region_grid.copy()
    region[5, :] = 2                       # a band right across the matrix
    state.region_grid = region
    flat = region.ravel()
    at_maps = {k: v.copy() for k, v in at_maps.items()}
    at_maps["Al"][flat == 2], at_maps["Si"][flat == 2] = 45.0, 55.0
    phase = np.where(flat > 0, 1, 0).astype(np.int32)
    state.phase_grid = phase.reshape(N_ROWS, N_COLS)
    state.locked_mask = None

    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    assert _row(tables.phases, phase_name="Al.cif")["n_particles"] > 1
    warn = [w for w in tables.warnings
            if w["code"] == "matrix_counted_as_particles"]
    assert warn, "the matrix now owns several components"
    assert "Al.cif" in warn[0]["message"]
    assert "not an inclusion count" in warn[0]["message"]


def test_the_coordinate_conventions_are_in_the_file_and_not_only_the_manual(fixture):
    c = _build(fixture).provenance["coordinates"]
    assert "0-based" in c["origin"]
    assert "row-major" in c["order"]
    assert "col * step_x_um" in c["x_um"]
    assert "DOWNWARD" in c["y_direction"]
    assert "INCLUSIVE" in c["bbox"]


def test_the_peak_overlap_sentence_is_about_this_scan_and_not_the_periodic_table(fixture):
    """Concrete, and true of the file it sits in."""
    state, at_maps = fixture
    at_maps = dict(at_maps)
    at_maps["Mn"] = np.full(N_PX, 1.0)
    at_maps["Fe"] = np.full(N_PX, 2.0)
    tables = X.build_tables(state, at_maps, geometry=_geom(),
                            options=X.ExportOptions(hash_source=False))
    block = tables.provenance["peak_overlaps"]
    pairs = {(d["interfering"], d["affected"]) for d in block["applies_to_this_scan"]}
    assert ("Mn", "Fe") in pairs
    text = next(d["note"] for d in block["applies_to_this_scan"]
                if (d["interfering"], d["affected"]) == ("Mn", "Fe"))
    assert "6.49 keV" in text and "6.40 keV" in text
    assert "OVER-REPORTS Fe" in text

    # A scan without manganese does not get the manganese sentence.
    plain = _build(fixture).provenance["peak_overlaps"]
    assert ("Mn", "Fe") not in {(d["interfering"], d["affected"])
                                for d in plain["applies_to_this_scan"]}
    assert "DECLARED" in plain["detection"]


def test_the_workbook_offers_filter_handles_and_labels_an_empty_sheet(tmp_path,
                                                                     fixture):
    """"sort phases.csv on worst_element_ratio" is this file's own advice."""
    from openpyxl import load_workbook

    tables = _build(fixture)
    res = X.write_export(tables, tmp_path,
                         options=X.ExportOptions(hash_source=False))
    wb = load_workbook(Path(res.folder) / "summary.xlsx")
    assert wb["phases"].auto_filter.ref
    assert wb["particles"].auto_filter.ref
    # definitions is empty here, and says which of "empty" and "broken" it is.
    assert wb["definitions"].auto_filter.ref is None
    assert "No region definitions were used" in str(wb["definitions"]["A2"].value)
